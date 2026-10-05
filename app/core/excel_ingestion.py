from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime, time, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pydantic import ValidationError

from app.config import settings
from app.schemas import FailureHistoryItem, TelemetryPoint, WorkHistoryItem

MEASUREMENT_FIELDS = {
    "water_cut_pct", "co2_pct", "chlorides_mg_l", "inhibitor_efficiency",
    "injection_deviation_pct", "corrosion_rate_mm_year", "wall_thickness_mm",
    "initial_wall_thickness_mm", "metal_loss_mm", "corrosion_load", "water_compatibility_issue",
}


class ExcelImportError(ValueError):
    pass


def load_profiles() -> dict:
    try:
        data = json.loads(settings.source_profiles_path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ExcelImportError(f"Source profile file is missing: {settings.source_profiles_path}") from exc
    except (json.JSONDecodeError, OSError) as exc:
        raise ExcelImportError(f"Source profile file is invalid: {exc}") from exc
    profiles = data.get("profiles")
    if not isinstance(profiles, dict):
        raise ExcelImportError("Source profile configuration must contain a profiles object")
    seen_source_ids: set[str] = set()
    for profile_id, profile in profiles.items():
        if not isinstance(profile, dict):
            raise ExcelImportError(f"Source profile '{profile_id}' must be an object")
        source_id = str(profile.get("source_id", profile_id)).strip().upper()
        if not source_id or len(source_id) > 128:
            raise ExcelImportError(f"Source profile '{profile_id}' has an invalid source_id")
        if source_id in seen_source_ids:
            raise ExcelImportError(f"More than one profile uses source_id '{source_id}'")
        seen_source_ids.add(source_id)
        if profile.get("kind") not in {"telemetry", "history", "json_schedule"}:
            raise ExcelImportError(f"Source profile '{profile_id}' has an unsupported kind")
        if profile.get("kind") != "json_schedule":
            if not isinstance(profile.get("field_map"), dict):
                raise ExcelImportError(f"Source profile '{profile_id}' must define field_map")
            if profile.get("sheet_name") and profile.get("sheet_names"):
                raise ExcelImportError(f"Source profile '{profile_id}' can set sheet_name or sheet_names, not both")
            if profile.get("sheet_name") is not None and not isinstance(profile["sheet_name"], str):
                raise ExcelImportError(f"sheet_name in profile '{profile_id}' must be a string or null")
            if profile.get("sheet_names") is not None and (
                not isinstance(profile["sheet_names"], list)
                or not profile["sheet_names"]
                or any(not isinstance(name, str) or not name.strip() for name in profile["sheet_names"])
                or len(set(profile["sheet_names"])) != len(profile["sheet_names"])
            ):
                raise ExcelImportError(f"sheet_names in profile '{profile_id}' must be a non-empty list of unique sheet names")
            if profile.get("sheet_header_rows") is not None and not isinstance(profile["sheet_header_rows"], dict):
                raise ExcelImportError(f"sheet_header_rows in profile '{profile_id}' must be an object")
            for field, aliases in profile["field_map"].items():
                if not isinstance(aliases, list) or any(not isinstance(alias, str) for alias in aliases):
                    raise ExcelImportError(f"field_map.{field} in profile '{profile_id}' must be a list of header aliases")
    return profiles


def _header_key(value: Any) -> str:
    return "".join(char for char in str(value or "").strip().casefold() if char.isalnum())


def _as_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, str):
        raw = value.strip()
        if not raw:
            raise ValueError("date is empty")
        try:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("date must be an Excel date or ISO 8601 value") from exc
    else:
        raise ValueError("date must be an Excel date or ISO 8601 value")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=ZoneInfo(settings.source_timezone))
    return parsed.astimezone(timezone.utc)


def _number(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    normalized = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    return float(normalized)


def _boolean(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    normalized = str(value).strip().casefold()
    if normalized in {"1", "true", "yes", "да", "есть", "обнаружена"}:
        return True
    if normalized in {"0", "false", "no", "нет", "отсутствует", "не обнаружена"}:
        return False
    raise ValueError("boolean value must be yes/no, true/false or 1/0")


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _row_values(headers: list[Any], values: tuple[Any, ...], field_map: dict) -> dict[str, Any]:
    columns = {_header_key(header): index for index, header in enumerate(headers) if header is not None}
    resolved = {}
    for canonical, aliases in field_map.items():
        candidates = [_header_key(canonical), *(_header_key(item) for item in aliases)]
        index = next((columns[key] for key in candidates if key in columns), None)
        if index is not None:
            resolved[canonical] = values[index] if index < len(values) else None
    return resolved


def _source_field_metadata(profile: dict, headers: list[Any]) -> dict[str, str]:
    columns = {_header_key(header): str(header) for header in headers if header is not None}
    resolved: dict[str, str] = {}
    for canonical, aliases in profile.get("field_map", {}).items():
        for key in [_header_key(canonical), *(_header_key(alias) for alias in aliases)]:
            if key in columns:
                resolved[canonical] = columns[key]
                break
    return resolved


def parse_excel(content: bytes, filename: str, profile_id: str, supplied_well_id: str | None = None) -> dict:
    if Path(filename).suffix.casefold() != ".xlsx":
        raise ExcelImportError("Only .xlsx files are supported. Save legacy .xls files as .xlsx first.")
    if len(content) > settings.max_excel_bytes:
        raise ExcelImportError(f"Workbook exceeds the {settings.max_excel_bytes // (1024 * 1024)} MB upload limit")
    profiles = load_profiles()
    profile = profiles.get(profile_id)
    if not isinstance(profile, dict):
        raise ExcelImportError(f"Unknown profile '{profile_id}'. Available profiles: {', '.join(sorted(profiles))}")
    kind = profile.get("kind")
    if kind not in {"telemetry", "history"}:
        raise ExcelImportError(f"Profile '{profile_id}' has unsupported kind '{kind}'")
    source_id = str(profile.get("source_id", profile_id)).upper()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelImportError("Could not read this workbook; confirm it is a valid .xlsx file") from exc
    try:
        explicit_names = profile.get("sheet_names")
        single_name = profile.get("sheet_name")
        explicit_selection = explicit_names is not None or single_name is not None
        if explicit_names is not None:
            missing_sheets = [name for name in explicit_names if name not in workbook.sheetnames]
            if missing_sheets:
                raise ExcelImportError(
                    f"Configured worksheet(s) not found: {', '.join(missing_sheets)}. "
                    f"Available sheets: {', '.join(workbook.sheetnames)}"
                )
            candidate_sheets = [workbook[name] for name in explicit_names]
        elif single_name:
            if single_name not in workbook.sheetnames:
                raise ExcelImportError(
                    f"Sheet '{single_name}' was not found; available sheets: {', '.join(workbook.sheetnames)}"
                )
            candidate_sheets = [workbook[single_name]]
        else:
            candidate_sheets = list(workbook.worksheets)

        selected_sheets: list[tuple[Any, int, list[Any], dict[str, str]]] = []
        skipped_sheets: list[str] = []
        source_columns_by_sheet: dict[str, dict[str, str]] = {}
        ignored_columns_by_sheet: dict[str, list[str]] = {}
        sheet_header_rows = profile.get("sheet_header_rows", {})
        for sheet in candidate_sheets:
            try:
                header_row_number = int(sheet_header_rows.get(sheet.title, profile.get("header_row", 1)))
            except (TypeError, ValueError) as exc:
                raise ExcelImportError(f"Invalid header row configured for sheet '{sheet.title}'") from exc
            if not 1 <= header_row_number <= settings.max_excel_rows:
                raise ExcelImportError(
                    f"Header row for sheet '{sheet.title}' must be between 1 and {settings.max_excel_rows}"
                )
            header_values = next(sheet.iter_rows(min_row=header_row_number, max_row=header_row_number,
                                                 values_only=True), None)
            headers = list(header_values or [])
            source_fields = _source_field_metadata(profile, headers)
            missing_headers = [field for field in profile.get("required_columns", []) if field not in source_fields]
            if "well_id" not in source_fields and not supplied_well_id:
                missing_headers.append("well_id (or supply the well_id form field)")
            if missing_headers:
                if explicit_selection:
                    recognized = ", ".join(f"{field}←{column}" for field, column in source_fields.items()) or "none"
                    raise ExcelImportError(
                        f"Sheet '{sheet.title}' is missing required columns: {', '.join(missing_headers)}. "
                        f"Recognized columns: {recognized}. Update config/source_profiles.json to map actual headers."
                    )
                skipped_sheets.append(sheet.title)
                continue
            selected_sheets.append((sheet, header_row_number, headers, source_fields))
            source_columns_by_sheet[sheet.title] = source_fields
            ignored_columns_by_sheet[sheet.title] = [
                str(header) for header in headers if header is not None and str(header) not in source_fields.values()
            ]

        if not selected_sheets:
            required = ", ".join(profile.get("required_columns", [])) or "profile fields"
            raise ExcelImportError(
                f"No worksheet matched profile '{profile_id}'. Required headers: {required}. "
                "Check header_row and field aliases in config/source_profiles.json."
            )

        telemetry_rows: dict[str, list[TelemetryPoint]] = defaultdict(list)
        history_work: dict[str, list[WorkHistoryItem]] = defaultdict(list)
        history_failures: dict[str, list[FailureHistoryItem]] = defaultdict(list)
        errors: list[dict] = []
        rows_read = accepted = 0
        work_values = {_header_key(item) for item in profile.get("work_event_values", [])}
        failure_values = {_header_key(item) for item in profile.get("failure_event_values", [])}
        try:
            scale_fields = {entry["field"]: float(entry.get("multiply", 1))
                            for entry in profile.get("scale_fields", [])
                            if isinstance(entry, dict) and entry.get("field")}
        except (TypeError, ValueError) as exc:
            raise ExcelImportError(f"Profile '{profile_id}' contains an invalid scale_fields multiplier") from exc

        for sheet, header_row_number, headers, _source_fields in selected_sheets:
            for excel_row, values in enumerate(
                sheet.iter_rows(min_row=header_row_number + 1,
                                max_row=header_row_number + settings.max_excel_rows + 1,
                                values_only=True),
                start=header_row_number + 1,
            ):
                if excel_row - header_row_number > settings.max_excel_rows:
                    if any(value is not None and str(value).strip() for value in values):
                        raise ExcelImportError(f"Workbook exceeds the {settings.max_excel_rows} row processing limit")
                    break
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                if rows_read >= settings.max_excel_rows:
                    raise ExcelImportError(f"Workbook exceeds the {settings.max_excel_rows} row processing limit")
                rows_read += 1
                row = _row_values(headers, values, profile.get("field_map", {}))
                well_id = _text(supplied_well_id) or _text(row.get("well_id"))
                if not well_id:
                    if len(errors) < 50:
                        errors.append({"sheet_name": sheet.title, "row": excel_row, "message": "well_id is blank"})
                    continue
                try:
                    if kind == "telemetry":
                        row["timestamp"] = _as_datetime(row.get("timestamp"))
                        normalized = {key: value for key, value in row.items()
                                      if key in MEASUREMENT_FIELDS or key == "timestamp"}
                        normalized["source_id"] = source_id
                        for field, multiplier in scale_fields.items():
                            if field in normalized and normalized[field] is not None:
                                normalized[field] = _number(normalized[field]) * multiplier
                        for field in MEASUREMENT_FIELDS - {"water_compatibility_issue"}:
                            if field in normalized:
                                normalized[field] = _number(normalized[field])
                        if "water_compatibility_issue" in normalized:
                            normalized["water_compatibility_issue"] = _boolean(normalized["water_compatibility_issue"])
                        point = TelemetryPoint.model_validate(normalized)
                        telemetry_rows[well_id].append(point)
                    else:
                        event_date = _as_datetime(row.get("event_date"))
                        event_raw = _text(row.get("event_type"))
                        event_key = _header_key(event_raw)
                        if not event_raw:
                            event_type = "failure" if _text(row.get("failure_type")) else "work"
                        elif event_key in failure_values:
                            event_type = "failure"
                        elif event_key in work_values:
                            event_type = "work"
                        else:
                            raise ValueError(f"unrecognized event_type '{event_raw}'; update event values in source profile")
                        if event_type == "failure":
                            failure_type = _text(row.get("failure_type")) or event_raw
                            if not failure_type:
                                raise ValueError("failure_type is required for a failure row")
                            history_failures[well_id].append(FailureHistoryItem(
                                failure_date=event_date, equipment=_text(row.get("equipment")),
                                failure_type=failure_type, cause=_text(row.get("cause")),
                                corrosion_detected=_boolean(row.get("corrosion_detected")),
                                repair=_text(row.get("repair")),
                                downtime_days=int(_number(row.get("downtime_days"))) if _number(row.get("downtime_days")) is not None else None,
                            ))
                        else:
                            work_type = _text(row.get("work_type")) or event_raw
                            if not work_type:
                                raise ValueError("work_type is required for a work row")
                            history_work[well_id].append(WorkHistoryItem(
                                date=event_date, work_type=work_type, reason=_text(row.get("reason")),
                                equipment=_text(row.get("equipment")), result=_text(row.get("result")),
                            ))
                    accepted += 1
                except (ValueError, TypeError, ValidationError) as exc:
                    if len(errors) < 50:
                        errors.append({"sheet_name": sheet.title, "row": excel_row,
                                       "well_id": well_id, "message": str(exc)[:500]})

        grouped: dict[str, dict] = {}
        all_wells = set(telemetry_rows) | set(history_work) | set(history_failures)
        for well_id in all_wells:
            grouped[well_id] = {
                "records": telemetry_rows.get(well_id, []),
                "work_history": history_work.get(well_id, []),
                "failure_history": history_failures.get(well_id, []),
            }
        file_hash = hashlib.sha256(content).hexdigest()
        return {
            "profile_id": profile_id, "source_id": source_id,
            "kind": kind, "sheet_name": selected_sheets[0][0].title,
            "sheet_names": [sheet.title for sheet, *_ in selected_sheets],
            "skipped_sheets": skipped_sheets,
            "source_columns": source_columns_by_sheet[selected_sheets[0][0].title],
            "source_columns_by_sheet": source_columns_by_sheet,
            "ignored_columns": list(dict.fromkeys(
                column for columns in ignored_columns_by_sheet.values() for column in columns
            )),
            "ignored_columns_by_sheet": ignored_columns_by_sheet,
            "rows_read": rows_read, "rows_accepted": accepted, "rows_rejected": rows_read - accepted,
            "errors": errors, "groups": grouped, "file_sha256": file_hash,
            "status": "ok" if accepted and not errors else ("partial" if accepted else "error"),
        }
    finally:
        workbook.close()
