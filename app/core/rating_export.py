"""Excel export of the computed well rating, in the approved layout.

Sheets: «Рейтинг» (one row per well, sorted by rank), «График ППР» (12 weeks from the Monday of
the export date), «Сводка», «Источники данных» (where each value came from) and «Параметры»
(the thresholds and deadlines the backend used). Every value is what GET /api/v1/rating returns,
written as a value rather than a formula, so the file matches the dashboard exactly.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core import rating

FONT = "Arial"
NAVY, NAVY2, BLUE = "0E2A47", "16375C", "0072CE"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_NOTE = Font(name=FONT, size=9, italic=True, color="5A6B7C")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BAD = Font(name=FONT, size=10, bold=True, color="B71C1C")
F_MUTED = Font(name=FONT, size=10, color="9AA7B4")
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_GROUP = PatternFill("solid", fgColor=NAVY2)
_THIN = Side(style="thin", color="DDE3EA")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Background / text colour per measure (same palette as the dashboard).
MEASURE_COLORS = {
    rating.MEASURE_REPAIR: ("FDECEA", "B71C1C"),
    rating.MEASURE_ENGINEER: ("FDEFE3", "BF4300"),
    rating.MEASURE_UD: ("FFF4D6", "8A5300"),
    rating.MEASURE_NONE: ("E8F5E9", "2E7D32"),
}

COLUMNS = [
    # key, header, width, group
    ("rank", "Место", 7, "Скважина"),
    ("well_id", "№ скважины", 11, "Скважина"),
    ("field", "Месторождение", 16, "Скважина"),
    ("equipment", "Спущенное оборудование", 22, "Скважина"),
    ("in_operation", "Статус", 12, "Скважина"),
    ("run_days", "Наработка, сут", 10, "Скважина"),
    ("liquid_rate_m3_day", "Жидкость, м³/сут", 10, "Скважина"),
    ("oil_l_day", "Нефть, л/сут", 10, "Скважина"),
    ("oil_pct", "Нефть, %", 8, "Скважина"),
    ("sampling", "Отбор проб", 10, "Среда и обработка"),
    ("corrosive_environment", "Среда", 15, "Среда и обработка"),
    ("tonnage_t", "Тоннаж, т", 10, "Среда и обработка"),
    ("noo_days", "НОО, сут", 10, "Среда и обработка"),
    ("demount_act", "Акт демонтажа", 11, "Демонтаж и разбор"),
    ("demount_corrosion", "Признак коррозии", 11, "Демонтаж и разбор"),
    ("base_act", "Акт разбора на базе", 11, "Демонтаж и разбор"),
    ("base_corrosion", "Результат разбора", 14, "Демонтаж и разбор"),
    ("category", "Категория скважины", 21, "Оценка"),
    ("protection", "Защита", 14, "Оценка"),
    ("model_ntc", "Модель НТЦ (0–1)", 10, "Оценка"),
    ("model_math", "Мат. модель (0–1)", 10, "Оценка"),
    ("index_raw", "Индекс без учёта защиты", 11, "Оценка"),
    ("index", "Индекс приоритета (0–100)", 11, "Оценка"),
    ("measure", "Мероприятие", 26, "Решение и график ППР"),
    ("action", "Предложение", 44, "Решение и график ППР"),
    ("era", "ЭРА ремонта", 11, "Решение и график ППР"),
    ("promine_status", "Документ Промайн", 16, "Решение и график ППР"),
    ("deadline", "Крайний срок", 12, "Решение и график ППР"),
    ("plan_date", "План ППР (ДДО)", 12, "Решение и график ППР"),
    ("plan_status", "Статус плана", 15, "Решение и график ППР"),
    ("data_gaps", "Полнота данных", 30, ""),
    ("comment", "Комментарий", 34, ""),
]
FIELD_LABELS = {key: header for key, header, *_ in COLUMNS}


def _yn(value: Any, yes: str, no: str) -> str:
    return "—" if value is None else (yes if value else no)


def _date(value: Any) -> date | None:
    if not value:
        return None
    return value if isinstance(value, date) else date.fromisoformat(str(value)[:10])


def _cell_values(well: dict) -> dict[str, Any]:
    measure = well.get("measure")
    return {
        "rank": well.get("rank"),
        "well_id": well["well_id"],
        "field": well.get("field"),
        "equipment": well.get("equipment"),
        "in_operation": _yn(well.get("in_operation"), "В работе", "Не в работе"),
        "run_days": well.get("run_days"),
        "liquid_rate_m3_day": well.get("liquid_rate_m3_day"),
        "oil_pct": well.get("oil_pct"),
        "oil_l_day": well.get("oil_l_day"),
        "action": well.get("action"),
        "data_gaps": ", ".join(well.get("data_gap_labels") or []) or "полные",
        "sampling": _yn(well.get("sampling"), "Да", "Нет"),
        "corrosive_environment": _yn(well.get("corrosive_environment"), "Коррозионная", "Некоррозионная"),
        "tonnage_t": well.get("tonnage_t"),
        "noo_days": well.get("noo_days"),
        "demount_act": _yn(well.get("demount_act"), "Есть", "Нет"),
        "demount_corrosion": _yn(well.get("demount_corrosion"), "Есть коррозия", "Коррозии нет") if well.get("demount_act") else "—",
        "base_act": _yn(well.get("base_act"), "Есть", "Нет"),
        "base_corrosion": (_yn(well.get("base_corrosion"), "Есть коррозия", "Коррозии нет")
                           if well.get("base_act") else "—"),
        "category": well.get("category_label") or "—",
        "protection": well.get("protection_label") or "Нет данных",
        "index_raw": well.get("index_raw"),
        "model_ntc": well.get("model_ntc"),
        "model_math": well.get("model_math"),
        "index": well.get("index"),
        "measure": well.get("measure_label") if measure else "Нет оценок моделей",
        "era": "Включить" if well.get("era_required") else "—",
        "promine_status": (well.get("promine_status_label") or "Не сформирован") if well.get("era_required")
        else (well.get("promine_status_label") or "—"),
        "deadline": _date(well.get("deadline")),
        "plan_date": _date(well.get("plan_date")),
        "plan_status": well.get("plan_status_label") or "—",
        "comment": well.get("comment") or well.get("skip_reason"),
    }


def _head(sheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row, column, label)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, CENTER, BORDER


def _rating_sheet(book: Workbook, wells: list[dict], generated: datetime) -> None:
    sheet = book.active
    sheet.title = "Рейтинг"
    sheet["A1"] = "Рейтинг скважин по риску коррозии"
    sheet["A1"].font = F_TITLE
    sheet["A2"] = (f"Выгрузка из системы «Контроль коррозии скважин» {generated:%d.%m.%Y %H:%M} UTC. "
                   "Значения рассчитаны по данным ДАРС, ХИМИК и ЭРА ремонта; источники — на листе "
                   "«Источники данных».")
    sheet["A2"].font = F_NOTE
    groups: list[list] = []
    for index, (_key, _header, _width, group) in enumerate(COLUMNS, start=1):
        if groups and groups[-1][0] == group:
            groups[-1][2] = index
        else:
            groups.append([group, index, index])
    for group, first, last in groups:
        sheet.cell(3, first, group)
        for column in range(first, last + 1):
            cell = sheet.cell(3, column)
            cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_GROUP, CENTER, BORDER
        if last > first:
            sheet.merge_cells(start_row=3, start_column=first, end_row=3, end_column=last)
    _head(sheet, 4, [header for _key, header, *_ in COLUMNS])
    for column, (_key, _header, width, _group) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[4].height = 44

    for offset, well in enumerate(wells):
        row = 5 + offset
        values = _cell_values(well)
        colors = MEASURE_COLORS.get(well.get("measure"))
        for column, (key, *_rest) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row, column, values[key])
            cell.font, cell.border = F_BASE, BORDER
            cell.alignment = LEFT if key in {"field", "equipment", "measure", "comment", "category", "action", "data_gaps"} else CENTER
            if key in {"model_ntc", "model_math"}:
                cell.number_format = "0.00"
            elif key in {"tonnage_t", "liquid_rate_m3_day", "oil_pct"}:
                cell.number_format = "0.0"
            elif key in {"noo_days", "index", "index_raw", "oil_l_day"}:
                cell.number_format = "0"
            elif key in {"deadline", "plan_date"}:
                cell.number_format = "DD.MM.YYYY"
            if key in {"index", "measure", "era"} and colors:
                cell.fill = PatternFill("solid", fgColor=colors[0])
                cell.font = Font(name=FONT, size=10, bold=key != "era", color=colors[1])
        for key, bad in (("corrosive_environment", "Коррозионная"), ("demount_corrosion", "Есть коррозия"),
                         ("base_corrosion", "Есть коррозия"), ("category", "Коррозионная"),
                         ("protection", "Недостаточна")):
            if values[key] == bad:
                sheet.cell(row, [k for k, *_ in COLUMNS].index(key) + 1).font = F_BAD
        status_cell = sheet.cell(row, [k for k, *_ in COLUMNS].index("plan_status") + 1)
        if well.get("plan_status") == "late":
            status_cell.fill, status_cell.font = PatternFill("solid", fgColor="FDECEA"), F_BAD
        elif well.get("plan_status") == "not_planned":
            status_cell.fill = PatternFill("solid", fgColor="FFF4D6")
            status_cell.font = Font(name=FONT, size=10, bold=True, color="8A5300")
    last = 4 + max(1, len(wells))
    sheet.freeze_panes = "C5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(COLUMNS))}{last}"
    sheet.sheet_view.zoomScale = 90
    if not wells:
        sheet["A5"] = "Рейтинг пуст: загрузите данные ДАРС, ХИМИК, ЭРА ремонта и выполните дневной цикл."
        sheet["A5"].font = F_NOTE


def _schedule_sheet(book: Workbook, wells: list[dict], today: date) -> None:
    sheet = book.create_sheet("График ППР")
    sheet["A1"] = "График ППР (ДДО) — мероприятия по рейтингу на 12 недель"
    sheet["A1"].font = F_TITLE
    sheet["A2"] = ("Код мероприятия стоит в неделе плановой даты; светлая заливка — окно до крайнего "
                   "срока; «!» — план позже срока. Скважины без мероприятий не показаны.")
    sheet["A2"].font = F_NOTE
    heads = ["Место", "№ скважины", "Месторождение", "Индекс", "Код", "Крайний срок", "План (ДДО)"]
    start = today - timedelta(days=today.weekday())
    weeks = [start + timedelta(days=7 * k) for k in range(12)]
    _head(sheet, 4, heads + [f"{week:%d.%m}" for week in weeks])
    for column in range(1, len(heads) + 1):
        sheet.cell(3, column).fill = FILL_HEAD
    for k, week in enumerate(weeks):
        cell = sheet.cell(3, len(heads) + 1 + k, f"Нед. {k + 1}")
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_GROUP, CENTER, BORDER
        if week <= today < week + timedelta(days=7):
            sheet.cell(4, len(heads) + 1 + k).fill = PatternFill("solid", fgColor=BLUE)
        sheet.column_dimensions[get_column_letter(len(heads) + 1 + k)].width = 8.5
    for column, width in enumerate((7, 11, 16, 8, 7, 12, 12), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[4].height = 30
    planned = [well for well in wells if well.get("measure") not in (None, rating.MEASURE_NONE)]
    for offset, well in enumerate(planned):
        row = 5 + offset
        deadline, plan = _date(well.get("deadline")), _date(well.get("plan_date"))
        code = well.get("measure_code") or ""
        values = [well.get("rank"), well["well_id"], well.get("field"), well.get("index"), code, deadline, plan]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.font, cell.border = F_BASE, BORDER
            cell.alignment = LEFT if column == 3 else CENTER
        sheet.cell(row, 6).number_format = sheet.cell(row, 7).number_format = "DD.MM.YYYY"
        background, foreground = MEASURE_COLORS[well["measure"]]
        for k, week in enumerate(weeks):
            cell = sheet.cell(row, len(heads) + 1 + k)
            cell.border, cell.alignment = BORDER, CENTER
            week_end = week + timedelta(days=6)
            if plan and week <= plan <= week_end:
                late = deadline is not None and plan > deadline
                cell.value = code + (" !" if late else "")
                cell.fill = PatternFill("solid", fgColor="1A2430" if late else foreground)
                cell.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
            elif deadline and week_end >= today and week <= deadline:
                cell.fill = PatternFill("solid", fgColor=background)
    if not planned:
        sheet["A5"] = "Мероприятий по рейтингу нет."
        sheet["A5"].font = F_NOTE
    sheet.freeze_panes = sheet.cell(5, len(heads) + 1)
    sheet.sheet_view.zoomScale = 90


def _summary_sheet(book: Workbook, wells: list[dict], counts: dict) -> None:
    sheet = book.create_sheet("Сводка")
    sheet["A1"] = "Сводка по рейтингу"
    sheet["A1"].font = F_TITLE

    def rows(start: int, title: list[str], data: list[list]) -> int:
        _head(sheet, start, title)
        for offset, values in enumerate(data, start=1):
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(start + offset, column, value)
                cell.font, cell.border = (F_BASE if column == 1 else F_BOLD), BORDER
                cell.alignment = LEFT if column == 1 else CENTER
        return start + len(data) + 2

    measures = [rating.MEASURE_REPAIR, rating.MEASURE_ENGINEER, rating.MEASURE_UD, rating.MEASURE_NONE]
    nxt = rows(3, ["Мероприятие", "Скважин"],
               [[rating.MEASURE_LABELS[key], counts.get(key, 0)] for key in measures]
               + [["Нет оценок моделей", counts.get("unscored", 0)], ["Всего в рейтинге", counts.get("total", 0)]])
    nxt = rows(nxt, ["Контроль исполнения", "Скважин"], [
        ["Требуют включения в ЭРА ремонта", counts.get(rating.MEASURE_REPAIR, 0)],
        ["…из них документ Промайн ещё не направлен", counts.get("promine_pending", 0)],
        ["План ППР позже крайнего срока", counts.get("plan_late", 0)],
        ["Мероприятие назначено, но не запланировано", counts.get("plan_missing", 0)],
    ])
    fields = sorted({well.get("field") or "не указано" for well in wells})
    nxt = rows(nxt, ["Месторождение", "ППР", "РИ", "УДР", "Всего"], [
        [name] + [sum(1 for well in wells if (well.get("field") or "не указано") == name and well.get("measure") == key)
                  for key in measures[:3]]
        + [sum(1 for well in wells if (well.get("field") or "не указано") == name)]
        for name in fields])
    rows(nxt, ["Категория скважины", "Скважин"], [
        [label, sum(1 for well in wells if well.get("category") == key)]
        for key, label in rating.CATEGORY_LABELS.items()])
    sheet.column_dimensions["A"].width = 46
    for letter in "BCDE":
        sheet.column_dimensions[letter].width = 14


def _sources_sheet(book: Workbook, wells: list[dict]) -> None:
    sheet = book.create_sheet("Источники данных")
    sheet["A1"] = "Откуда взяты значения рейтинга"
    sheet["A1"].font = F_TITLE
    _head(sheet, 3, ["№ скважины", "Поле", "Источник", "Дата данных", "Пояснение"])
    row = 4
    for well in wells:
        provenance = well.get("provenance") or {}
        for key, item in provenance.items():
            if not isinstance(item, dict):
                continue
            when = item.get("date")
            values = [well["well_id"], FIELD_LABELS.get(key, key), item.get("source") or "расчёт системы",
                      datetime.fromisoformat(when).replace(tzinfo=None) if when else None, item.get("note")]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row, column, value)
                cell.font, cell.border = F_BASE, BORDER
                cell.alignment = LEFT
            sheet.cell(row, 4).number_format = "DD.MM.YYYY HH:MM"
            row += 1
        if well.get("skip_reason"):
            for column, value in enumerate([well["well_id"], "Модели не запускались", "проверка готовности",
                                            None, well["skip_reason"]], start=1):
                cell = sheet.cell(row, column, value)
                cell.font, cell.border, cell.alignment = F_MUTED, BORDER, LEFT
            row += 1
    for letter, width in zip("ABCDE", (12, 24, 22, 17, 80)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A4"


def _parameters_sheet(book: Workbook, payload: dict) -> None:
    sheet = book.create_sheet("Параметры")
    sheet["A1"] = "Параметры рейтинга (настройки системы на момент выгрузки)"
    sheet["A1"].font = F_TITLE
    settings_block = payload.get("settings", {})
    sheet["A3"] = "Объединение моделей НТЦ и мат."
    sheet["C3"] = "Максимум" if settings_block.get("combine") == "max" else "Среднее"
    sheet["A4"] = "Горизонт прогноза моделей, сут"
    sheet["C4"] = settings_block.get("horizon_days")
    for cell in ("A3", "A4"):
        sheet[cell].font = F_BASE
    _head(sheet, 6, ["Мероприятие", "Код", "Индекс от", "Индекс до", "Срок, дней", "ЭРА ремонта"])
    for offset, item in enumerate(settings_block.get("measures", []), start=7):
        for column, value in enumerate([item["label"], item["code"], item["index_from"], item["index_to"],
                                        item.get("deadline_days"), "Включить" if item.get("era_required") else "—"],
                                       start=1):
            cell = sheet.cell(offset, column, value)
            cell.font, cell.border = F_BASE, BORDER
            cell.alignment = LEFT if column == 1 else CENTER
    protection = settings_block.get("protection") or {}
    sheet["A13"] = (f"Защита (режим {protection.get('mode', '—')}): эффективна, если ингибитор ≥ "
                    f"{protection.get('min_inhibitor', '—')} при отклонении закачки < {protection.get('max_injection_dev', '—')} % "
                    f"или НКТ в защитном исполнении, и это подтверждено скоростью коррозии < {protection.get('max_rate', '—')} мм/год "
                    f"без коррозии в последнем акте. Тогда баллы среды учитываются с коэффициентом {protection.get('factor', '—')}.")
    sheet["A13"].font = F_NOTE
    sheet["A12"] = ("Индекс = ОКРУГЛ(объединение оценок НТЦ и мат. модели × 100). Срок мероприятия "
                    "отсчитывается от даты, когда скважине впервые назначено текущее мероприятие. "
                    "Пороги и сроки задаются переменными MAI_RATING_*.")
    sheet["A12"].font = F_NOTE
    for letter, width in zip("ABCDEF", (40, 8, 11, 11, 11, 13)):
        sheet.column_dimensions[letter].width = width


def build_workbook(payload: dict) -> bytes:
    """Workbook bytes for the GET /rating payload (wells already ranked and evaluated)."""
    generated = datetime.now(timezone.utc)
    wells = payload.get("wells", [])
    book = Workbook()
    _rating_sheet(book, wells, generated)
    _schedule_sheet(book, wells, generated.date())
    _summary_sheet(book, wells, payload.get("counts", {}))
    _sources_sheet(book, wells)
    _parameters_sheet(book, payload)
    for sheet in book.worksheets:
        sheet.sheet_properties.tabColor = {"Рейтинг": NAVY, "График ППР": BLUE}.get(sheet.title, "8FA3B8")
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    book.active = 0
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
