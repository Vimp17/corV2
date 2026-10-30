"""Excel 97-2003 workbooks (.xls): converted to .xlsx in memory, so every reader of the system
(source profiles, documents, the НКТ registry) handles them as it handles .xlsx.

Cell values keep their type: dates become datetimes, whole numbers become integers (a well number
7781 stays 7781, not 7781.0), errors become empty cells; merged ranges are kept.
"""
from __future__ import annotations

from io import BytesIO

# OLE2 compound file: .xls (and .doc, .ppt — told apart by the workbook stream xlrd looks for).
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def is_legacy_workbook(content: bytes) -> bool:
    return content[:8] == OLE2_MAGIC


def xls_to_xlsx(content: bytes) -> bytes:
    """The .xlsx bytes of an .xls workbook; ValueError when the file is not a readable .xls."""
    import xlrd
    from openpyxl import Workbook

    try:
        book = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except Exception:  # noqa: BLE001 — formatting records of some writers are not readable
        try:
            book = xlrd.open_workbook(file_contents=content)
        except Exception as exc:  # noqa: BLE001
            raise ValueError("Не удалось прочитать книгу Excel 97-2003 (.xls).") from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in book.sheets():
        target = workbook.create_sheet(title=(sheet.name or f"Лист{len(workbook.worksheets) + 1}")[:31])
        for row in range(sheet.nrows):
            for column in range(sheet.ncols):
                value = _value(sheet.cell(row, column), book.datemode)
                if value is not None:
                    target.cell(row=row + 1, column=column + 1, value=value)
        for row_low, row_high, column_low, column_high in getattr(sheet, "merged_cells", []) or []:
            if row_high - row_low > 1 or column_high - column_low > 1:
                target.merge_cells(start_row=row_low + 1, start_column=column_low + 1,
                                   end_row=row_high, end_column=column_high)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _value(cell, datemode: int):
    import xlrd

    kind, value = cell.ctype, cell.value
    if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if kind == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(value, datemode)
        except (ValueError, OverflowError):
            return None
    if kind == xlrd.XL_CELL_BOOLEAN:
        return bool(value)
    if kind == xlrd.XL_CELL_NUMBER:
        return int(value) if float(value).is_integer() else value
    text = str(value)
    return text if text.strip() else None
