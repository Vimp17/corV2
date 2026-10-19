"""Text of a document: Excel cells, the PDF text layer, or OCR of page images via LiteLLM.

OCR uses the vision Chat Completions interface: each page is sent as an image in a user
message to MAI_OCR_MODEL, which returns the page text (tables as Markdown).
"""
from __future__ import annotations

import base64
import logging
from dataclasses import dataclass, field
from io import BytesIO

from app.config import settings
from app.core import llm

logger = logging.getLogger(__name__)

OCR_SYSTEM_PROMPT = (
    "Ты — модуль OCR. Перепиши весь текст со страницы документа дословно и в исходном порядке: "
    "заголовки, реквизиты, таблицы (в формате Markdown), подписи и рукописные пометки, если они "
    "читаются. Ничего не добавляй от себя, не исправляй, не переводи и не пересказывай. "
    "Нечитаемое место обозначь [неразборчиво]. Если на странице нет текста, верни пустой ответ."
)


class TextUnavailable(RuntimeError):
    """The text cannot be obtained now (OCR not configured or failing); retrying later may help."""


@dataclass
class TextResult:
    text: str
    method: str                      # text_layer / ocr / mixed / excel_table / image_ocr
    pages: int | None = None
    notes: list[str] = field(default_factory=list)


def ocr_configured() -> bool:
    return bool(settings.ocr_url and settings.ocr_model)


def _data_url(image) -> str:
    from PIL import Image

    image = image.convert("RGB")
    limit = settings.ocr_max_side_px
    if max(image.size) > limit:
        image.thumbnail((limit, limit), Image.LANCZOS)
    buffer = BytesIO()
    image.save(buffer, format="JPEG", quality=90)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def ocr_image(image, label: str) -> str:
    """Text of one page image from the OCR model."""
    if not ocr_configured():
        raise TextUnavailable("OCR-модель не настроена (MAI_OCR_MODEL).")
    messages = [
        {"role": "system", "content": OCR_SYSTEM_PROMPT},
        {"role": "user", "content": [
            {"type": "text", "text": f"{label}. Перепиши текст страницы."},
            {"type": "image_url", "image_url": {"url": _data_url(image)}},
        ]},
    ]
    try:
        return llm.chat(messages, model=settings.ocr_model, url=settings.ocr_url,
                        api_key=settings.ocr_key, timeout=settings.ocr_timeout_seconds,
                        url_variable="MAI_OCR_API_URL")
    except llm.LLMError as exc:
        raise TextUnavailable(f"OCR-модель недоступна: {exc}") from exc


def pdf_text(content: bytes, filename: str) -> TextResult:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(content)
    except Exception as exc:  # pdfium raises PdfiumError for damaged or encrypted files
        raise ValueError("PDF повреждён или защищён паролем.") from exc
    try:
        count = len(pdf)
        if count > settings.docs_max_pages:
            raise ValueError(f"В PDF {count} стр., обрабатывается не больше {settings.docs_max_pages} "
                             "(MAI_DOCS_MAX_PAGES).")
        parts: list[str] = []
        methods: set[str] = set()
        notes: list[str] = []
        for index in range(count):
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                layer = textpage.get_text_range().replace("\r\n", "\n").strip()
                textpage.close()
                if len(layer) >= settings.docs_min_page_chars:
                    parts.append(f"=== Страница {index + 1} ===\n{layer}")
                    methods.add("text_layer")
                    continue
                width, height = page.get_size()
                scale = min(300 / 72, settings.ocr_max_side_px / max(width, height, 1))
                image = page.render(scale=scale).to_pil()
                recognised = ocr_image(image, f"Страница {index + 1} из {count} документа «{filename}»")
                parts.append(f"=== Страница {index + 1} (OCR) ===\n{recognised.strip()}")
                methods.add("ocr")
            finally:
                page.close()
        method = "mixed" if len(methods) > 1 else (methods.pop() if methods else "text_layer")
        return TextResult(text="\n\n".join(parts).strip(), method=method, pages=count, notes=notes)
    finally:
        pdf.close()


def image_text(content: bytes, filename: str) -> TextResult:
    from PIL import Image, ImageSequence

    try:
        image = Image.open(BytesIO(content))
        frames = [frame.copy() for _, frame in zip(range(settings.docs_max_pages + 1), ImageSequence.Iterator(image))]
    except Exception as exc:
        raise ValueError("Изображение повреждено или не читается.") from exc
    if len(frames) > settings.docs_max_pages:
        raise ValueError(f"В файле больше {settings.docs_max_pages} страниц (MAI_DOCS_MAX_PAGES).")
    parts = []
    for index, frame in enumerate(frames):
        recognised = ocr_image(frame, f"Страница {index + 1} из {len(frames)} документа «{filename}»")
        parts.append(f"=== Страница {index + 1} (OCR) ===\n{recognised.strip()}")
    return TextResult(text="\n\n".join(parts).strip(), method="image_ocr", pages=len(frames))


def workbook_text(content: bytes) -> TextResult:
    """Cells of every sheet as lines "a | b | c" (an act or protocol kept as an Excel form)."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Не удалось прочитать книгу Excel.") from exc
    try:
        parts: list[str] = []
        size = 0
        for sheet in workbook.worksheets:
            lines = []
            for values in sheet.iter_rows(max_row=settings.max_excel_rows, values_only=True):
                cells = []
                for value in values:
                    if value is None or not str(value).strip():
                        continue
                    if hasattr(value, "strftime"):
                        value = value.strftime("%d.%m.%Y")
                    cells.append(str(value).strip())
                if cells:
                    line = " | ".join(cells)
                    lines.append(line)
                    size += len(line)
                if size > settings.docs_max_text_chars * 2:
                    break
            if lines:
                parts.append(f"=== Лист «{sheet.title}» ===\n" + "\n".join(lines))
        return TextResult(text="\n\n".join(parts).strip(), method="excel_table", pages=len(workbook.sheetnames))
    finally:
        workbook.close()


def detect_profile(content: bytes, filename: str) -> str | None:
    """The source profile (ДАРС / ХИМИК / ЭРА) whose table layout the workbook has, if any.

    The header row is looked for in the first 30 rows (a table may sit under a title). A profile
    matches when its parser accepts at least half of the rows and recognises at least three
    columns. Among matches wins the one recognising the most columns and ignoring the
    fewest; then the one recognising columns no other match knows (дебит, статус → ДАРС;
    тоннаж, НОО → ХИМИК); a remaining tie — a table of chemistry only (CO₂, хлориды, pH…) —
    goes to the laboratory profile, so it counts as a sample."""
    from app.core.excel_ingestion import ExcelImportError, load_profiles, parse_excel

    name = filename if filename.casefold().endswith(".xlsx") else "document.xlsx"
    candidates: list[dict] = []
    for profile_id, profile in load_profiles().items():
        if profile.get("kind") not in {"telemetry", "history"}:
            continue
        try:
            parsed = parse_excel(content, name, profile_id, auto_header=True)
        except (ExcelImportError, ValueError):
            continue
        accepted, read = parsed.get("rows_accepted") or 0, parsed.get("rows_read") or 0
        fields = {field for columns in (parsed.get("source_columns_by_sheet") or {}).values() for field in columns}
        recognised = sum(len(columns) for columns in (parsed.get("source_columns_by_sheet") or {}).values())
        if not accepted or accepted * 2 < read or recognised < 3:
            continue
        candidates.append({"profile_id": profile_id, "fields": fields, "recognised": recognised,
                           "unknown": len(parsed.get("ignored_columns") or []) - len(parsed.get("extra_columns") or []),
                           "ignored": len(parsed.get("ignored_columns") or []), "accepted": accepted,
                           "lab": bool(profile.get("lab_samples"))})
    if not candidates:
        return None
    for item in candidates:
        others = set().union(*(other["fields"] for other in candidates if other is not item)) if len(candidates) > 1 else set()
        item["exclusive"] = len(item["fields"] - others)
    best = max(candidates, key=lambda item: (item["recognised"] - item["ignored"], item["exclusive"], item["lab"],
                                             item["accepted"]))
    return best["profile_id"]


def document_text(document: dict, content: bytes) -> TextResult:
    kind = document["file_kind"]
    if kind == "pdf":
        return pdf_text(content, document["filename"])
    if kind == "image":
        return image_text(content, document["filename"])
    if kind == "xlsx":
        return workbook_text(content)
    raise ValueError(f"Неизвестный вид файла: {kind}")
