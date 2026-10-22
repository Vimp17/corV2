"""Реестр коррозии НКТ: an Excel table with one row per well (or per failure), read without AI.

The table is recognised by its header: a well column and at least four registry columns
(факты выявления коррозии НКТ, ННО, режим работы ЭЦН, признак новизны НКТ, группа прочности,
состав подвески, марка стали, тип повреждения…). A two-row header is understood: «Состав
подвески» over «верх | середина | низ», and a row of column numbers under the header is skipped.

For every accepted row:
  - месторождение, средняя обводнённость, Cl⁻ (хлориды), CO₂ % мол. and средний дебит жидкости
    go to the system fields as a sample of the well (source DOC_NKT) and count in the rating
    (corrosive environment). The sample is dated by the date above the table («по состоянию на
    01.10.2026»), else by the day the file was received;
  - the row as it is in the file, text columns included (ДО, куст, пласт, режим ЭЦН, состав
    подвески, марка стали, завод, тип повреждения…), is stored as the well's registry entry:
    shown in the well card and given to the AI agent, not used by the rating.
Corrosion facts and failure counts carry no dates, so they are not turned into ERA events.

A well not yet known to the system is accepted: the registry is itself a list of wells. A well
number met on two different fields is ambiguous and its rows are not loaded.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app import db
from app.config import settings
from app.core.documents.excel_reader import DATE_TEXT, _cell, _moment, _norm, _number, _well_text
from app.core.documents.extract import EARLIEST_DATE, clean_well_id, match_well
from app.schemas import TelemetryPoint

SOURCE_ID = "DOC_NKT"
TITLE = "Реестр коррозии НКТ"
DOC_TYPE = "excel:nkt_registry"
MAX_ROWS = 20000
MAX_COLUMNS = 80
HEADER_SCAN_ROWS = 30
MIN_MARKERS = 4
MAX_STORED_RECORDS = 100
MAX_REJECTED_STORED = 50

# (key, label, kind, system field, words, formula, unit)
#   kind     text / number / any (a number when the cell is numeric, else text) / skip;
#   words    regex on the normalised header («средняя обводненность, %»);
#   formula  regex on the header without spaces in Latin letters («so42-», «co2,%мол.»);
#   unit     regex on the normalised header that must match too (CO₂ in мл/г or in % мол.).
# The first matching column wins; specific columns come before general ones.
COLUMNS: list[tuple[str, str, str, str | None, str | None, str | None, str | None]] = [
    ("row_no", "№", "skip", None, r"^(№|n|no|#)\.?$|^№?\s*п\s*/?\s*п\.?$", None, None),
    ("subsidiary", "ДО", "text", None, r"^до\.?$|дочерн", None, None),
    ("field", "Месторождение", "text", "field", r"месторожден", None, None),
    ("pad", "Куст", "any", None, r"^куст", None, None),
    ("pump_depth_m", "Глубина спуска насоса, м", "number", None, r"глубин\w*\s+спуск", None, None),
    ("failure_depth_m", "Глубина отказа от устья, м", "number", None,
     r"глубин.*(отказ|поврежд)|глубин\w*\s+от\s+усть", None, None),
    ("formation", "Пласт", "text", None, r"^пласт(\s|,|$)", None, None),
    ("corrosion_facts", "Факты выявления коррозии НКТ", "any", None, r"факт\w*.*коррози", None, None),
    ("nno_days", "ННО, сут", "number", None, r"(^|[^а-яa-z])нно([^а-яa-z]|$)|наработк\w*\s+на\s+отказ", None, None),
    ("esp_mode", "Режим работы ЭЦН", "text", None, r"режим", None, None),
    ("tubing_new", "Признак новизны НКТ", "text", None, r"новизн", None, None),
    ("tubing_diameter", "Диаметр НКТ", "any", None, r"диаметр", None, None),
    ("string_top", "Состав подвески: верх", "text", None, r"(подвеск|материал).*верх", None, None),
    ("string_middle", "Состав подвески: середина", "text", None, r"(подвеск|материал).*серед", None, None),
    ("string_bottom", "Состав подвески: низ", "text", None, r"(подвеск|материал).*низ", None, None),
    ("damaged_strength_group", "Группа прочности повреждённой НКТ", "text", None, r"прочност", None, None),
    ("wellhead_temp_c", "Температура на устье, °C", "number", None, r"температур.*усть", None, None),
    ("reservoir_temp_c", "Температура пластовая, °C", "number", None,
     r"температур.*пласт|пластов\w*\s+температур", None, None),
    ("wellhead_pressure_atm", "Давление на устье, атм", "number", None, r"давлен.*усть", None, None),
    ("liquid_rate_m3_day", "Средний дебит жидкости, м³/сут", "number", "liquid_rate_m3_day",
     r"дебит\w*\s+жидк|^(средн\w*\s+)?дебит", None, None),
    ("water_cut_pct", "Средняя обводнённость, %", "number", "water_cut_pct", r"обводн", None, None),
    ("water_density", "Плотность воды, кг/м³", "number", None, r"плотност", None, None),
    ("mineralization_mg_l", "Общая минерализация, мг/л", "number", None, r"минерализ", None, None),
    ("ph", "pH воды", "number", None, r"(^|[^a-z])ph([^a-z]|$)", None, None),
    ("kvch_mg_l", "КВЧ, мг/л", "number", None, r"квч|взвешен", None, None),
    ("gas_factor", "Газовый фактор, м³/т", "number", None, r"газов\w*\s+фактор", None, None),
    ("steel_grade", "Марка стали", "text", None, r"марк\w*\s+стал", None, None),
    ("plant", "Завод", "text", None, r"^завод|изготовит", None, None),
    ("failures_count", "Отказы за весь период, кол-во", "number", None, r"отказ", None, None),
    ("damage_type", "Тип повреждения", "text", None, r"(тип|вид)\w*\s+повреж", None, None),
    ("wall_thickness_mm", "Толщина стенки повреждённой НКТ, мм", "number", None, r"толщин", None, None),
    ("ca", "Ca²⁺", "number", None, r"кальци", r"^ca(\d|\+|$)", None),
    ("mg", "Mg²⁺", "number", None, r"магни", r"^mg(\d|\+|$)", None),
    ("na_k", "Na⁺ + K⁺", "number", None, r"натри", r"^na(\+|k|$)", None),
    ("fe", "Fe общ.", "number", None, r"желез", r"^fe", None),
    ("cl", "Cl⁻ (хлориды), мг/л", "number", "chlorides_mg_l", r"хлорид", r"^cl(-|−|$)", None),
    ("so4", "SO₄²⁻", "number", None, r"сульфат", r"^so4", None),
    ("hco3", "HCO₃⁻", "number", None, r"гидрокарбонат", r"^hco3", None),
    ("co3", "CO₃²⁻", "number", None, r"^карбонат", r"^co3", None),
    ("co2_ml_g", "CO₂, мл/г", "number", None, None, r"^co2", r"мл"),
    ("co2_mol_pct", "CO₂, % мол.", "number", "co2_pct", None, r"^co2", r"%|мол"),
    ("h2s_mg_l", "H₂S, мг/л", "number", None, None, r"^h2s", r"мг"),
    ("h2s_mol_pct", "H₂S, % мол.", "number", None, None, r"^h2s", r"%|мол"),
]
LABELS = {key: label for key, label, *_ in COLUMNS} | {"well_id": "Скважина"}
SYSTEM_FIELDS = {key: system for key, _, _, system, *_ in COLUMNS if system}
# Columns that tell the registry apart from other tables with a well column.
MARKERS = {"corrosion_facts", "nno_days", "esp_mode", "tubing_new", "damaged_strength_group", "string_top",
           "string_middle", "string_bottom", "steel_grade", "failures_count", "damage_type", "pump_depth_m",
           "failure_depth_m", "wall_thickness_mm", "subsidiary", "suspension"}
SUSPENSION = re.compile(r"состав\w*\s+подвеск")
# Columns describing where the well is, not what was measured: a row needs more than these.
LOCATION = {"subsidiary", "field", "pad", "formation"}
WELL_HEADER = re.compile(r"^(№\s*)?(скв(ажин[аы]?)?\.?)(\s*№)?$|^№\s*скв|^номер\s+скв|^скв(ажин[аы]?)?[\s.,:;(/№-]")
# Latin letters typed in place of Cyrillic look-alikes («Cкважина» with a Latin C): read as Cyrillic.
_CYRILLIC = str.maketrans({"a": "а", "c": "с", "e": "е", "o": "о", "p": "р", "x": "х", "k": "к", "y": "у",
                           "b": "в", "h": "н", "m": "м", "t": "т"})
REGISTRY_DATE = re.compile(r"состояни\w*\s+на|на\s+дату|дата|период|^на\s+\d", re.IGNORECASE)
_LATIN = str.maketrans({"а": "a", "с": "c", "о": "o", "е": "e", "к": "k", "н": "h", "м": "m", "р": "p", "х": "x"})


def _is_well_header(text: str) -> bool:
    norm = _norm(text)
    return bool(WELL_HEADER.search(norm) or WELL_HEADER.search(norm.translate(_CYRILLIC)))


def _formula(text: str) -> str:
    return "".join(text.split()).translate(_LATIN)


def column_key(label: str, used: set[str]) -> str | None:
    """Registry column of a header, None for a column the registry does not know."""
    text = _norm(label)
    compact = _formula(text)
    for key, _label, _kind, _system, words, formula, unit in COLUMNS:
        if key in used:
            continue
        if words and re.search(words, text) or formula and re.search(formula, compact) and (
                unit is None or re.search(unit, text)):
            return key
    return None


def _labels(rows: list[list[Any]], r: int) -> tuple[dict[int, str], int]:
    """Column labels of the header at row r (1-based) and the first row after it.

    A second header row (sub-headers under a merged cell: «Состав подвески» over «верх | середина
    | низ») is joined to the label above it; a row of column numbers (1, 2, 3…) is skipped."""
    top = rows[r - 1]
    labels = {c: " ".join(str(v).split()) for c, v in enumerate(top) if isinstance(v, str) and v.strip()}
    nxt = r
    if r < len(rows):
        below = rows[r]
        subs = {c: " ".join(str(v).split()) for c, v in enumerate(below) if isinstance(v, str) and v.strip()}
        numbers = [v for v in below if isinstance(v, (int, float)) and not isinstance(v, bool)]
        # Sub-headers sit under a merged cell, so some of them have nothing written above.
        if len(subs) >= 2 and len(numbers) <= 1 and any(c not in labels for c in subs):
            parent = ""
            for c in range(max(len(top), len(below))):
                if c in labels:
                    parent = labels[c]
                if c in subs:
                    labels[c] = f"{labels[c]} {subs[c]}" if c in labels else f"{parent} {subs[c]}".strip()
            nxt = r + 1
    if nxt < len(rows):
        values = [v for v in rows[nxt] if v is not None and str(v).strip()]
        if len(values) >= 3 and all(isinstance(v, int) and not isinstance(v, bool) for v in values) \
                and all(b - a == 1 for a, b in zip(values, values[1:])):
            nxt += 1
    return labels, nxt


def find_header(rows: list[list[Any]]) -> dict | None:
    """{"row", "data_from", "well", "columns": {col: key}, "titles": {col: label}} or None."""
    for r in range(1, min(len(rows), HEADER_SCAN_ROWS) + 1):
        labels, data_from = _labels(rows, r)
        well = next((c for c, text in labels.items() if _is_well_header(text)), None)
        if well is None:
            continue
        used: set[str] = set()
        columns: dict[int, str] = {}
        markers = {"suspension"} if any(SUSPENSION.search(_norm(text)) for text in labels.values()) else set()
        for c, text in sorted(labels.items()):
            if c == well:
                continue
            key = column_key(text, used)
            if key:
                used.add(key)
                columns[c] = key
                markers |= {key} & MARKERS
        if len(markers) >= MIN_MARKERS:
            return {"row": r, "data_from": data_from + 1, "well": well, "columns": columns, "titles": labels}
    return None


def _read_rows(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        return [(sheet.title, [list(values)[:MAX_COLUMNS] for values in
                               sheet.iter_rows(max_row=MAX_ROWS, values_only=True)])
                for sheet in workbook.worksheets]
    finally:
        workbook.close()


def detect(content: bytes) -> bool:
    """True when a sheet of the workbook has the registry header."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return False
    try:
        for sheet in workbook.worksheets:
            rows = [list(values)[:MAX_COLUMNS] for values in
                    sheet.iter_rows(max_row=HEADER_SCAN_ROWS + 2, values_only=True)]
            if find_header(rows):
                return True
    finally:
        workbook.close()
    return False


def _registry_date(title: str, rows: list[list[Any]], header_row: int) -> tuple[date, str] | None:
    """The date written above the table: «Реестр … по состоянию на 01.10.2026», «Дата: …»."""
    for r in range(1, header_row):
        row = rows[r - 1]
        for c, value in enumerate(row):
            if isinstance(value, (datetime, date)):
                moment = value if isinstance(value, datetime) else datetime.combine(value, time.min)
                return moment.date(), _cell(title, r, c, "Дата", value)
            if isinstance(value, str) and DATE_TEXT.search(value) and (REGISTRY_DATE.search(value) or r == 1):
                moment = _moment(value)
                if moment is not None:
                    local = moment.astimezone(ZoneInfo(settings.source_timezone)).date()
                    return local, f"Лист «{title}», {get_column_letter(c + 1)}{r}: {value.strip()}"[:300]
    return None


def _value(raw: Any, kind: str) -> Any:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if kind == "number":
        return _number(raw)
    if isinstance(raw, (datetime, date)):
        return raw.strftime("%d.%m.%Y")
    if kind == "any":
        number = _number(raw)
        if number is not None:
            return int(number) if number.is_integer() else number
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return " ".join(str(raw).split())[:300]


def read_workbook(content: bytes, known_wells: list[str], received: date) -> dict:
    """Registry rows (accepted and not), the registry date and which columns were used."""
    sheets = _read_rows(content)
    kinds = {key: kind for key, _, kind, *_ in COLUMNS}
    records: list[dict] = []
    labels = dict(LABELS)
    used_columns: dict[str, str] = {}
    other_columns: list[str] = []
    unused_columns: list[str] = []
    registry_date: tuple[date, str] | None = None
    for title, rows in sheets:
        header = find_header(rows)
        if header is None:
            continue
        registry_date = registry_date or _registry_date(title, rows, header["row"])
        titles = header["titles"]
        columns = dict(header["columns"])
        for c, text in titles.items():
            if c == header["well"] or c in columns:
                continue
            key = f"col:{get_column_letter(c + 1)}"
            columns[c] = key
            labels[key] = text[:100]
            kinds[key] = "any"
        for c, key in columns.items():
            if kinds.get(key) == "skip":
                continue
            if key in SYSTEM_FIELDS:
                used_columns[key] = titles[c]
        empty = 0
        filled: set[int] = set()
        for r in range(header["data_from"], len(rows) + 1):
            row = rows[r - 1]
            well_raw = _well_text(row[header["well"]]) if header["well"] < len(row) else ""
            if not well_raw:
                empty += 1
                if empty > 20:
                    break
                continue
            empty = 0
            if _norm(well_raw) in {"итого", "всего"}:
                continue
            fields = {"well_id": {"value": well_raw, "status": "ok",
                                  "quote": _cell(title, r, header["well"], titles[header["well"]], row[header["well"]])}}
            for c, key in columns.items():
                if kinds[key] == "skip" or c >= len(row):
                    continue
                value = _value(row[c], kinds[key])
                if value is None:
                    continue
                filled.add(c)
                fields[key] = {"value": value, "quote": _cell(title, r, c, titles[c], row[c]), "status": "ok"}
            records.append({"well_raw": well_raw, "fields": fields, "sheet": title, "row": r})
        for c, key in columns.items():
            if kinds[key] == "skip" or key in SYSTEM_FIELDS:
                continue
            (other_columns if c in filled else unused_columns).append(titles[c])

    _check(records, known_wells)
    accepted = [record for record in records if record["accepted"]]
    rejected = [record for record in records if not record["accepted"]]
    when, quote = registry_date or (received, None)
    date_note = None
    if not EARLIEST_DATE <= when <= date.today() + timedelta(days=1):
        date_note = f"Дата над таблицей {when:%d.%m.%Y} неправдоподобна — принята дата получения файла."
        when, quote = received, None
    elif registry_date is None:
        date_note = f"Дата в файле не указана — показатели датированы днём получения файла ({received:%d.%m.%Y})."
    return {"records": accepted, "rejected": rejected, "labels": labels, "date": when, "date_quote": quote,
            "date_note": date_note, "columns": {"system": list(dict.fromkeys(used_columns.values())),
                                                "other": list(dict.fromkeys(other_columns)),
                                                "unused": list(dict.fromkeys(unused_columns))}}


def _check(records: list[dict], known: list[str]) -> None:
    """Match each row to a well; a number met on two fields is ambiguous and not loaded."""
    fields_of: dict[str, set[str]] = defaultdict(set)
    for record in records:
        issues: list[str] = []
        well = match_well(record["well_raw"], known)
        record["new_well"] = False
        if well is None:
            cleaned = clean_well_id(record["well_raw"])
            if cleaned and re.search(r"\d", cleaned):
                well, record["new_well"] = cleaned, True
            else:
                issues.append(f"Не распознан номер скважины «{record['well_raw']}».")
        record["well_id"] = well
        values = [key for key in record["fields"] if key != "well_id" and key not in LOCATION]
        if well and not values:
            issues.append("В строке нет данных, кроме расположения скважины.")
        field = (record["fields"].get("field") or {}).get("value")
        if well and field:
            fields_of[well].add(_norm(field))
        record["issues"] = issues
    for record in records:
        well = record["well_id"]
        if well and len(fields_of.get(well, ())) > 1:
            record["issues"].append(f"Скважина {well} есть на разных месторождениях — номер неоднозначен, строка не загружена.")
        record["accepted"] = not record["issues"]


def _stored(record: dict) -> dict:
    """A record as kept in the document's extraction (shown in the documents section)."""
    return {"well_id": record["well_id"], "accepted": record["accepted"], "issues": record["issues"],
            "new_well": record.get("new_well", False), "fields": record["fields"]}


def load(document: dict, result: dict) -> dict:
    """Store the accepted rows: a sample per well (system fields) and the registry entry."""
    from app.core.documents.load import batch_id

    document_id = document["document_id"]
    db.delete_document_data(document_id)
    moment = datetime.combine(result["date"], time(12, 0), tzinfo=timezone.utc)
    by_well: dict[str, list[dict]] = defaultdict(list)
    for record in result["records"]:
        by_well[record["well_id"]].append(record)
    samples = 0
    for well_id, rows in by_well.items():
        # The latest row in the file gives the sample; every row is kept in the registry entry.
        values = {}
        for record in rows:
            for key, system in SYSTEM_FIELDS.items():
                value = (record["fields"].get(key) or {}).get("value")
                if value is not None:
                    values[system] = str(value)[:256] if system == "field" else value
        valid = {}
        for name, value in values.items():
            try:  # an out-of-range value (e.g. a negative rate) is dropped, the others are kept
                TelemetryPoint.model_validate({"timestamp": moment, name: value})
                valid[name] = value
            except ValueError:
                continue
        if valid:
            point = TelemetryPoint.model_validate({"timestamp": moment, "source_id": SOURCE_ID, **valid})
            db.save_batch(well_id, batch_id(document_id, well_id), [point])
            samples += 1
        db.save_source_payload(SOURCE_ID, well_id, {
            "document_id": document_id, "filename": document["filename"], "date": result["date"].isoformat(),
            "rows": [{key: field["value"] for key, field in record["fields"].items()} for record in rows],
            "labels": {key: result["labels"].get(key, key) for record in rows for key in record["fields"]}})
        db.record_source_status(SOURCE_ID, well_id, "ok", f"{TITLE}: «{document['filename']}»", len(rows),
                                origin="document")
    return {"wells": sorted(by_well), "records": len(result["records"]), "samples": samples, "source_id": SOURCE_ID,
            "new_wells": sorted({record["well_id"] for record in result["records"] if record.get("new_well")}),
            "source_columns": {name: name for name in result["columns"]["system"]},
            "extra_columns": result["columns"]["other"], "unused_columns": result["columns"]["unused"]}


def extraction(result: dict) -> dict:
    """What is kept in the document for the documents section (the first rows only)."""
    accepted, rejected = result["records"], result["rejected"]
    return {"type_method": "rules", "records_total": len(accepted) + len(rejected), "records_accepted": len(accepted),
            "records": [_stored(record) for record in accepted[:MAX_STORED_RECORDS]]
            + [_stored(record) for record in rejected[:MAX_REJECTED_STORED]],
            "records_shown_limit": MAX_STORED_RECORDS, "field_labels": result["labels"],
            "date_note": result["date_note"]}


def not_loaded_reason(result: dict) -> str:
    if not result["records"] and not result["rejected"]:
        return "Шапка реестра найдена, но строк со скважинами нет."
    issues = list(dict.fromkeys(issue for record in result["rejected"] for issue in record["issues"]))
    return "Данные не загружены: " + " ".join(issues[:4])


def well_registry(well_id: str) -> dict | None:
    """The well's registry entry for the well card and the AI agent: {date, filename, rows: [[label, value]…]}."""
    payload = db.load_source_payloads(well_id).get(SOURCE_ID)
    if not isinstance(payload, dict) or not payload.get("rows"):
        return None
    labels = payload.get("labels") or {}
    order = [key for key, *_ in COLUMNS]
    rows = []
    for row in payload["rows"]:
        keys = sorted((key for key in row if key != "well_id"),
                      key=lambda key: (order.index(key) if key in order else len(order), key))
        rows.append([[labels.get(key) or LABELS.get(key, key), row[key]] for key in keys])
    return {"title": TITLE, "date": payload.get("date"), "filename": payload.get("filename"),
            "document_id": payload.get("document_id"), "rows": rows}
