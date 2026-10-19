"""Reading measurements from an arbitrary Excel workbook without AI.

Used for Excel documents that are not a ДАРС / ХИМИК / ЭРА export. Each sheet is read as
  a table   - a header row (in the first 30 rows) with a well column; indicator columns are
              recognised by the source-profile aliases and by name patterns (CO₂, хлориды,
              обводнённость…); the date comes from a date column or from the sheet title
              («Дата отбора: 16.10.2026»);
  a form    - one sample per sheet: «Скважина № 1234», «Дата отбора 16.10.2026» and rows
              «Хлориды, мг/л | 21000», «pH | 6,8».
A record is accepted only with a well from the register, a date and at least one indicator.
Indicators without a system field (pH, hardness…) become additional indicators. Every value
keeps the sheet and cell it came from, shown on the dashboard in place of a quote.
"""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.config import settings
from app.core.documents.extract import EARLIEST_DATE, match_well

MAX_ROWS = 2000
MAX_COLUMNS = 60
HEADER_SCAN_ROWS = 30
MAX_REJECTED_STORED = 50

# Recognised indicators: (field, name pattern, label). Specific patterns go first.
FIELD_PATTERNS: list[tuple[str, str, str]] = [
    ("initial_wall_thickness_mm", r"(начальн|исходн)\w*\s+толщин", "Начальная толщина стенки, мм"),
    ("wall_thickness_mm", r"толщин\w*(\s+стенк)?", "Толщина стенки, мм"),
    ("metal_loss_mm", r"потер\w*\s+металл", "Потеря металла, мм"),
    ("corrosion_rate_mm_year", r"скорост\w*\s+коррози", "Скорость коррозии, мм/год"),
    ("inhibitor_efficiency", r"эффективн\w*\s+(ингибитор|защит)", "Эффективность ингибитора, доля"),
    ("injection_deviation_pct", r"отклонен\w*\s+(закачк|дозировк|подач)", "Отклонение закачки, %"),
    ("water_cut_pct", r"обводн", "Обводнённость, %"),
    ("chlorides_mg_l", r"хлорид|(^|[^a-zа-я])cl([^a-zа-я]|$)", "Хлориды, мг/л"),
    ("co2_pct", r"(^|[^a-zа-я0-9])(co2|со2|co₂|со₂)([^a-zа-я0-9]|$)|углекисл|диоксид\w*\s+углерод", "CO₂, %"),
    ("liquid_rate_m3_day", r"дебит\w*\s+жидк|(^|[^a-zа-я])qж", "Дебит жидкости, м³/сут"),
    ("tonnage_t", r"тоннаж", "Тоннаж, т"),
    ("noo_days", r"(^|[^a-zа-я])ноо([^a-zа-я]|$)", "НОО, сут"),
]
FIELD_LABELS = {field: label for field, _, label in FIELD_PATTERNS} | {
    "well_id": "Скважина", "sample_date": "Дата пробы / замера", "field": "Месторождение",
    "extra_measurements": "Другие показатели",
}
WELL_LABEL = re.compile(r"скважин|(^|[^a-zа-я])скв([^a-zа-я]|$)|^well", re.IGNORECASE)
WELL_INLINE = re.compile(r"(?:скважин[аыуе]?|скв\.?)\s*(?:№|n|#|no\.?)?\s*:?\s*([A-Za-zА-Яа-яЁё0-9][\w\-/.]*)",
                         re.IGNORECASE)
DATE_LABEL = re.compile(r"дата|date|^число$", re.IGNORECASE)
DATE_TEXT = re.compile(r"(\d{1,2}\.\d{1,2}\.\d{2,4}(?:\s+\d{1,2}:\d{2})?|\d{4}-\d{2}-\d{2})")
# The date of a sample is preferred to the date of the protocol.
DATE_PRIORITY = ("отбор", "проб", "анализ", "замер", "измерен", "протокол", "")
FIELD_LABEL = re.compile(r"месторожден", re.IGNORECASE)
IDENTIFIER = re.compile(r"(^|[^a-zа-я])(№\s*п/?п|п/?п|куст|цех|участок|код|пласт|бригад|примечан|номер|№\s*пробы|"
                        r"протокол|лаборатор|лаборант|исполнит|подпис|телефон|итого|всего)", re.IGNORECASE)


def _norm(text: Any) -> str:
    return " ".join(str(text or "").replace("ё", "е").replace("Ё", "Е").split()).casefold()


def _alias_map() -> dict[str, str]:
    """Header aliases of the telemetry source profiles: normalised header -> field."""
    from app.core.excel_ingestion import _header_key, load_profiles

    aliases: dict[str, str] = {}
    try:
        profiles = load_profiles()
    except Exception:
        return aliases
    for profile in profiles.values():
        if profile.get("kind") != "telemetry":
            continue
        for field, names in (profile.get("field_map") or {}).items():
            for name in [field, *names]:
                aliases.setdefault(_header_key(name), field)
    return aliases


def field_for(label: Any, aliases: dict[str, str]) -> str | None:
    """System field of a column or row label, or None for an unknown indicator."""
    from app.core.excel_ingestion import _header_key

    text = _norm(label)
    if not text:
        return None
    field = aliases.get(_header_key(text))
    if field in {"timestamp", "well_id", "in_operation", "water_compatibility_issue"}:
        return None
    if field:
        return field
    if FIELD_LABEL.search(text):
        return "field"
    for name, pattern, _label in FIELD_PATTERNS:
        if re.search(pattern, text):
            return name
    return None


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, (bool, datetime, date)):
        return None
    if isinstance(value, (int, float)):
        return float(value) if value == value and abs(value) != float("inf") else None
    text = str(value).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not re.fullmatch(r"[-+]?\d+(\.\d+)?", text):
        return None
    return float(text)


def _well_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def _moment(value: Any) -> datetime | None:
    from app.core.excel_ingestion import _as_datetime

    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str):
        match = DATE_TEXT.search(value)
        if not match:
            return None
        value = match.group(1)
    try:
        return _as_datetime(value)
    except (TypeError, ValueError):
        return None


def _cell(sheet: str, row: int, column: int, label: str, value: Any) -> str:
    shown = value.strftime("%d.%m.%Y") if isinstance(value, (datetime, date)) else value
    return f"Лист «{sheet}», {get_column_letter(column + 1)}{row}: {label} = {shown}"[:300]


def _entry(value: Any, quote: str) -> dict:
    return {"value": value, "quote": quote, "status": "ok"}


def _adjust(field: str, label: str, value: float) -> float:
    # Inhibitor efficiency in percent («85 %») is stored as a fraction.
    if field == "inhibitor_efficiency" and (value > 1.5 or "%" in label):
        return round(value / 100, 4)
    return value


class _Sheet:
    def __init__(self, title: str, rows: list[list[Any]]) -> None:
        self.title = title
        self.rows = rows

    def texts(self, limit: int | None = None):
        for r, row in enumerate(self.rows if limit is None else self.rows[:limit], start=1):
            for c, value in enumerate(row):
                if isinstance(value, str) and value.strip():
                    yield r, c, value.strip()

    def right_of(self, r: int, c: int) -> tuple[int, Any] | None:
        row = self.rows[r - 1]
        for index in range(c + 1, len(row)):
            if row[index] is not None and str(row[index]).strip():
                return index, row[index]
        return None


def _sheet_date(sheet: _Sheet, limit: int | None = None) -> tuple[datetime, str] | None:
    """The date of the sheet: «Дата отбора: 16.10.2026» or a «Дата …» label with a date beside it."""
    found: list[tuple[int, datetime, str]] = []
    for r, c, text in sheet.texts(limit):
        if not DATE_LABEL.search(text):
            continue
        low = _norm(text)
        rank = next(index for index, word in enumerate(DATE_PRIORITY) if word in low)
        moment = _moment(text)
        quote = f"Лист «{sheet.title}», {get_column_letter(c + 1)}{r}: {text}"[:300]
        if moment is None:
            beside = sheet.right_of(r, c)
            if beside is not None:
                moment = _moment(beside[1])
                quote = _cell(sheet.title, r, beside[0], text, beside[1])
        if moment is not None:
            found.append((rank, moment, quote))
    if not found:
        return None
    found.sort(key=lambda item: item[0])
    return found[0][1], found[0][2]


def _header_row(sheet: _Sheet, aliases: dict[str, str]) -> tuple[int, int, int | None, dict[int, str], dict[int, str]] | None:
    """(row, well column, date column, field columns, other columns) of a table, or None."""
    for r, row in enumerate(sheet.rows[:HEADER_SCAN_ROWS], start=1):
        labels = {c: str(value).strip() for c, value in enumerate(row) if isinstance(value, str) and value.strip()}
        well_column = next((c for c, text in labels.items()
                            if WELL_LABEL.search(_norm(text)) and not re.search(r"\d", text) and len(text) <= 40), None)
        if well_column is None:
            continue
        date_column = next((c for c, text in labels.items() if c != well_column and DATE_LABEL.search(_norm(text))), None)
        fields, others = {}, {}
        for c, text in labels.items():
            if c in {well_column, date_column}:
                continue
            field = field_for(text, aliases)
            if field and field not in fields.values():
                fields[c] = field
            elif not IDENTIFIER.search(_norm(text)):
                others[c] = text
        if fields or others:
            return r, well_column, date_column, fields, others
    return None


def _record(well_raw: str, well_quote: str, moment: datetime | None, date_quote: str | None,
            values: dict[str, tuple[Any, str]], extras: list[dict], known: list[str]) -> dict:
    if moment is not None:
        # The day as written in the file (source time zone), not the UTC day of local midnight.
        moment = moment.astimezone(ZoneInfo(settings.source_timezone))
    fields: dict[str, dict] = {"well_id": _entry(well_raw, well_quote)}
    issues: list[str] = []
    well = match_well(well_raw, known)
    if well is None:
        issues.append(f"Скважины «{well_raw}» нет в справочнике.")
    if moment is not None and not EARLIEST_DATE <= moment.date() <= date.today() + timedelta(days=1):
        issues.append(f"Неправдоподобная дата {moment:%d.%m.%Y}.")
        moment = None
        fields["sample_date"] = {"value": None, "quote": date_quote, "status": "implausible"}
    elif moment is None:
        issues.append("Не найдена дата (столбец «Дата» или «Дата отбора: …» в шапке).")
        fields["sample_date"] = {"value": None, "quote": None, "status": "missing"}
    else:
        fields["sample_date"] = _entry(moment.date().isoformat(), date_quote)
    for field, (value, quote) in values.items():
        fields[field] = _entry(value, quote)
    fields["extra_measurements"] = ({"value": extras, "quote": None, "status": "ok"} if extras
                                    else {"value": None, "quote": None, "status": "missing"})
    if not values and not extras:
        issues.append("Нет показателей.")
    return {"well_id": well, "fields": fields, "accepted": not issues, "issues": issues,
            "moment": moment.isoformat() if moment else None}


def _read_table(sheet: _Sheet, header: tuple, aliases: dict[str, str], known: list[str]) -> list[dict]:
    header_row, well_column, date_column, fields, others = header
    titles = {c: str(v).strip() for c, v in enumerate(sheet.rows[header_row - 1]) if v is not None}
    sheet_date = _sheet_date(sheet, header_row - 1) if date_column is None else None
    records, empty = [], 0
    for r in range(header_row + 1, len(sheet.rows) + 1):
        row = sheet.rows[r - 1]
        well_raw = _well_text(row[well_column]) if well_column < len(row) else ""
        if not well_raw:
            empty += 1
            if empty > 20:
                break
            continue
        empty = 0
        if IDENTIFIER.search(_norm(well_raw)) or _norm(well_raw) in {"итого", "всего"}:
            continue
        moment, date_quote = None, None
        if date_column is not None and date_column < len(row):
            moment = _moment(row[date_column])
            date_quote = _cell(sheet.title, r, date_column, titles.get(date_column, "Дата"), row[date_column])
        elif sheet_date is not None:
            moment, date_quote = sheet_date
        values: dict[str, tuple[Any, str]] = {}
        for c, field in fields.items():
            raw = row[c] if c < len(row) else None
            if field == "field":
                if isinstance(raw, str) and raw.strip():
                    values[field] = (raw.strip()[:256], _cell(sheet.title, r, c, titles[c], raw))
                continue
            number = _number(raw)
            if number is not None:
                values[field] = (_adjust(field, titles[c], number), _cell(sheet.title, r, c, titles[c], raw))
        extras = []
        for c, name in others.items():
            number = _number(row[c] if c < len(row) else None)
            if number is not None:
                extras.append({"name": " ".join(name.split())[:100], "value": number,
                               "quote": _cell(sheet.title, r, c, name, row[c])})
        records.append(_record(well_raw, _cell(sheet.title, r, well_column, titles[well_column], row[well_column]),
                               moment, date_quote, values, extras, known))
    return records


def _read_form(sheet: _Sheet, aliases: dict[str, str], known: list[str]) -> list[dict]:
    """One sample per sheet: the well, the date and «indicator | value» rows."""
    well_raw, well_quote = None, None
    for r, c, text in sheet.texts():
        match = WELL_INLINE.search(text)
        if match and re.search(r"\d", match.group(1)):
            well_raw, well_quote = match.group(1).rstrip(".,;"), _cell(sheet.title, r, c, "Скважина", text)
            break
        if WELL_LABEL.search(_norm(text)) and not re.search(r"\d", text):
            beside = sheet.right_of(r, c)
            if beside is not None and re.search(r"\d", _well_text(beside[1])):
                well_raw, well_quote = _well_text(beside[1]), _cell(sheet.title, r, beside[0], text, beside[1])
                break
    if well_raw is None:
        return []
    sheet_date = _sheet_date(sheet)
    values: dict[str, tuple[Any, str]] = {}
    extras: list[dict] = []
    for r, row in enumerate(sheet.rows, start=1):
        label_index = next((c for c, value in enumerate(row) if isinstance(value, str) and re.search(r"[A-Za-zА-Яа-яЁё]", value)), None)
        if label_index is None:
            continue
        label = str(row[label_index]).strip()
        low = _norm(label)
        if (len(label) > 80 or WELL_LABEL.search(low) or DATE_LABEL.search(low) or IDENTIFIER.search(low)
                or WELL_INLINE.search(label)):
            continue
        unit, number, column = None, None, None
        for c in range(label_index + 1, len(row)):
            candidate = row[c]
            if candidate is None or not str(candidate).strip():
                continue
            number = _number(candidate)
            if number is not None:
                column = c
                break
            if unit is None and isinstance(candidate, str) and len(candidate) <= 15:
                unit = candidate.strip()
            else:
                break
        if number is None:
            continue
        name = f"{label}, {unit}" if unit and unit.casefold() not in low else label
        field = field_for(name, aliases)
        quote = _cell(sheet.title, r, column, name, row[column])
        if field == "field":
            continue
        if field and field not in values:
            values[field] = (_adjust(field, name, number), quote)
        elif not field:
            extras.append({"name": " ".join(name.split())[:100], "value": number, "quote": quote})
    moment, date_quote = sheet_date if sheet_date else (None, None)
    return [_record(well_raw, well_quote, moment, date_quote, values, extras, known)]


def read_workbook(content: bytes, known_wells: list[str]) -> dict:
    """Records found in the workbook (accepted and not) and a summary of what was found."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    aliases = _alias_map()
    records: list[dict] = []
    layouts: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [list(values)[:MAX_COLUMNS] for values in worksheet.iter_rows(max_row=MAX_ROWS, values_only=True)]
            sheet = _Sheet(worksheet.title, rows)
            header = _header_row(sheet, aliases)
            if header is not None:
                found = _read_table(sheet, header, aliases, known_wells)
                if found:
                    layouts.append("table")
                    records.extend(found)
                    continue
            found = _read_form(sheet, aliases, known_wells)
            if found:
                layouts.append("form")
                records.extend(found)
    finally:
        workbook.close()
    accepted = [record for record in records if record["accepted"]]
    rejected = [record for record in records if not record["accepted"]]
    return {"records": accepted + rejected[:MAX_REJECTED_STORED], "accepted": len(accepted), "rejected": len(rejected),
            "layout": "form" if layouts and set(layouts) == {"form"} else "table" if layouts else None,
            "field_labels": FIELD_LABELS}


def not_found_reason(result: dict) -> str:
    """Why nothing was loaded, in words for the user."""
    records = result["records"]
    if not records:
        return ("В файле не найдена скважина: нужен столбец «Скважина» (или «№ скв.») в таблице "
                "или строка «Скважина № …» в форме.")
    issues = list(dict.fromkeys(issue for record in records for issue in record["issues"]))
    return "Данные не загружены: " + " ".join(issues[:4])
