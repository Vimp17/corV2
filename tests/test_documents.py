"""Document pipeline: checks without a database, then the full flow on PostgreSQL.

The full flow uploads the synthetic ДАРС / ХИМИК / ЭРА workbooks and the demo documents of
scripts/build_demo_documents.py, with LiteLLM replaced by tests.fake_litellm.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import time
from datetime import date, timedelta

import pytest

from tests.conftest import ROOT, TEST_TMP
from tests.fake_litellm import FakeLiteLLM

from app.core.documents import extract, intake

TODAY = date.today()

BASE_SPEC = {
    "title": "Акт разбора на базе", "target": "work_event", "work_type": "Акт разбора на базе",
    "source_id": "DOC_BASE",
    "fields": {
        "well_id": {"label": "Номер скважины", "type": "string", "required": True},
        "act_date": {"label": "Дата разбора", "type": "date", "required": True},
        "corrosion_detected": {"label": "Обнаружена коррозия", "type": "boolean"},
        "corrosion_type": {"label": "Вид коррозии", "type": "string"},
    },
    "informative_fields": ["corrosion_detected", "corrosion_type"],
}
TEXT = ("АКТ РАЗБОРА ОБОРУДОВАНИЯ НА БАЗЕ\nДата разбора: 12.03.2026\n| Скважина № | DEMO-002 |\n"
        "Коррозия обнаружена: да\nВид коррозии: язвенная")


def _raw(**overrides):
    record = {
        "well_id": {"value": "DEMO-002", "quote": "Скважина № | DEMO-002"},
        "act_date": {"value": "2026-03-12", "quote": "Дата разбора: 12.03.2026"},
        "corrosion_detected": {"value": True, "quote": "Коррозия обнаружена: да"},
        "corrosion_type": {"value": "язвенная", "quote": "Вид коррозии: язвенная"},
    }
    record.update(overrides)
    return record


# --------------------------------------------------------------------------- no database
def test_quote_checks_tolerate_tables_and_spacing():
    text = extract.normalize(TEXT)
    assert extract.quote_found("Скважина №  DEMO-002", text)
    assert extract.quote_found("дата  разбора: 12.03.2026", text)
    assert not extract.quote_found("Скважина № DEMO-777 после ремонта", text)


def test_check_record_accepts_confirmed_values():
    record = extract.check_record(_raw(), BASE_SPEC, extract.normalize(TEXT), ["DEMO-001", "DEMO-002"],
                                  today=date(2026, 10, 19))
    assert record["accepted"], record["issues"]
    assert record["well_id"] == "DEMO-002"
    assert extract.values_of(record)["act_date"] == "2026-03-12"


def test_value_without_quote_in_text_is_dropped():
    raw = _raw(corrosion_type={"value": "общая", "quote": "Вид коррозии: общая"},
               corrosion_detected={"value": True, "quote": None})
    record = extract.check_record(raw, BASE_SPEC, extract.normalize(TEXT), ["DEMO-002"], today=date(2026, 10, 19))
    assert record["fields"]["corrosion_type"]["status"] == "quote_not_found"
    assert record["fields"]["corrosion_detected"]["status"] == "no_quote"
    assert not record["accepted"]          # no informative field left


def test_unknown_well_and_implausible_date_reject_the_record():
    raw = _raw(well_id={"value": "DEMO-999", "quote": "DEMO-002"})
    record = extract.check_record(raw, BASE_SPEC, extract.normalize(TEXT), ["DEMO-002"], today=date(2026, 10, 19))
    assert not record["accepted"]           # the quote does not support DEMO-999
    record = extract.check_record(_raw(), BASE_SPEC, extract.normalize(TEXT), ["DEMO-002"], today=date(2025, 1, 1))
    assert not record["accepted"]
    assert any("Неправдоподобная дата" in issue for issue in record["issues"])


def test_well_matching():
    known = ["DEMO-002", "1234Б", "77"]
    assert extract.match_well("скв. № demo-002", known) == "DEMO-002"
    assert extract.match_well("1234б", known) == "1234Б"
    assert extract.match_well("077", known) == "77"
    assert extract.match_well("5555", known) is None
    assert extract.match_well("5555", []) == "5555"


def test_percent_written_as_fraction_agrees_with_quote():
    assert extract.value_agrees(0.85, "number", "Эффективность ингибитора: 85 %", "inhibitor_efficiency")
    assert extract.value_agrees(18500.0, "number", "Хлориды | 18 500 мг/л", "chlorides_mg_l")
    assert not extract.value_agrees(0.5, "number", "Эффективность ингибитора: 85 %", "inhibitor_efficiency")


def test_keywords_classify_demo_texts():
    from app.core.documents.types import load_document_types
    from scripts.build_demo_documents import base_act_lines, memo_lines, passport_lines, scan_lines

    types = load_document_types()
    assert extract.classify_by_keywords("\n".join(base_act_lines(TODAY)), types) == "base_inspection_act"
    assert extract.classify_by_keywords("\n".join(scan_lines(TODAY)), types) == "demount_act"
    assert extract.classify_by_keywords("\n".join(passport_lines(TODAY)), types) == "well_passport"
    assert extract.classify_by_keywords("\n".join(memo_lines(TODAY)), types) is None


def test_file_kinds():
    assert intake.detect_kind(b"%PDF-1.7 ...", "a.pdf") == "pdf"
    assert intake.detect_kind(b"\x89PNG\r\n\x1a\n....", "scan.png") == "image"
    with pytest.raises(intake.DocumentRejected):
        intake.detect_kind(b"\xd0\xcf\x11\xe0....", "old.xls")
    with pytest.raises(intake.DocumentRejected):
        intake.detect_kind(b"PK\x03\x04....", "letter.docx")


# --------------------------------------------------------------------------- PostgreSQL
needs_db = pytest.mark.skipif(not os.environ.get("MAI_TEST_DATABASE_URL"),
                              reason="set MAI_TEST_DATABASE_URL to run the PostgreSQL tests")


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    out = tmp_path_factory.mktemp("data")
    subprocess.run([sys.executable, str(ROOT / "scripts" / "build_sample_data.py"), "--end-date", TODAY.isoformat(),
                    "--out", str(out)], check=True, capture_output=True)
    subprocess.run([sys.executable, str(ROOT / "scripts" / "build_demo_documents.py"), "--end-date",
                    TODAY.isoformat(), "--out", str(out / "documents")], check=True, capture_output=True)
    return out


@pytest.fixture()
def fake(monkeypatch):
    from app.core import llm
    from scripts.build_demo_documents import scan_lines

    stub = FakeLiteLLM("\n".join(scan_lines(TODAY)))
    monkeypatch.setattr(llm, "chat", stub.chat)
    return stub


@pytest.fixture()
def client(generated, fake):
    from fastapi.testclient import TestClient
    from sqlalchemy import text

    from app import db
    from app.main import app

    with TestClient(app) as test_client:
        with db.engine().begin() as conn:
            conn.execute(text("TRUNCATE documents, telemetry, ingestion_batches, history_events, history_batches, "
                              "source_health, source_documents, well_rating, analysis_runs, era_proposals CASCADE"))
        for profile, name in (("dars", "sample_dars.xlsx"), ("khimik", "sample_khimik.xlsx"), ("era", "sample_era.xlsx")):
            with open(generated / name, "rb") as handle:
                response = test_client.post("/api/v1/ingestion/excel", data={"profile_id": profile},
                                            files={"file": (name, handle)})
            assert response.status_code == 201, response.text
        yield test_client


def _upload(client, path, wait=True):
    with open(path, "rb") as handle:
        response = client.post(f"/api/v1/documents?wait={'true' if wait else 'false'}",
                               files=[("files", (path.name, handle))])
    assert response.status_code == 201, response.text
    return response.json()["documents"][0]


def _well(client, well_id):
    response = client.get(f"/api/v1/rating/wells/{well_id}")
    assert response.status_code == 200, response.text
    return response.json()


def _gap_keys(well):
    return {item["key"] for item in well["data_gaps"]}


@needs_db
def test_base_act_pdf_closes_the_missing_act_gap(client, generated):
    assert "base_act" in _gap_keys(_well(client, "DEMO-002"))
    document = _upload(client, generated / "documents" / "akt_razbora_DEMO-002.pdf")
    assert document["status"] == "loaded", document["status_message"]
    assert document["doc_type"] == "base_inspection_act"
    assert document["text_method"] == "text_layer"
    assert document["well_ids"] == ["DEMO-002"]
    well = _well(client, "DEMO-002")
    assert "base_act" not in _gap_keys(well)
    assert well["base_act"] is True and well["base_corrosion"] is True
    assert well["category"] == "corrosive"
    assert well["provenance"]["base_act"]["source"] == "Акт разбора на базе (документ)"
    assert [item["document_id"] for item in well["documents"]] == [document["document_id"]]

    # The same file again is recognised by content; reprocessing loads nothing twice.
    again = _upload(client, generated / "documents" / "akt_razbora_DEMO-002.pdf")
    assert again["duplicate"] is True and again["document_id"] == document["document_id"]
    response = client.post(f"/api/v1/documents/{document['document_id']}/reprocess?wait=true")
    assert response.json()["status"] == "loaded"
    from app import db
    acts = [event for event in db.load_history_events("DEMO-002") if event["source_id"] == "DOC_BASE"]
    assert len(acts) == 1

    # Deleting the document removes its data: the gap is back.
    assert client.delete(f"/api/v1/documents/{document['document_id']}").status_code == 200
    assert "base_act" in _gap_keys(_well(client, "DEMO-002"))


@needs_db
def test_scanned_act_goes_through_ocr(client, generated, fake):
    document = _upload(client, generated / "documents" / "akt_demontazha_DEMO-005_scan.jpg")
    assert document["status"] == "loaded", document["status_message"]
    assert document["text_method"] == "image_ocr"
    assert any(call["kind"] == "ocr" and call["model"] == "fake-ocr" for call in fake.calls)
    well = _well(client, "DEMO-005")
    assert well["demount_act"] is True and well["demount_corrosion"] is True
    assert "base_act" in _gap_keys(well)          # a fresh demount without a base act


@needs_db
def test_lab_protocol_excel_form_counts_as_samples(client, generated):
    assert "samples" in _gap_keys(_well(client, "DEMO-012"))
    document = _upload(client, generated / "documents" / "protokol_khim_analiza_15.xlsx")
    assert document["status"] == "loaded", document["status_message"]
    # The table under the protocol title is found by its header row: a ХИМИК import, no AI.
    assert document["doc_type"] == "source:khimik"
    assert document["well_ids"] == ["DEMO-012", "DEMO-013"]
    well = _well(client, "DEMO-012")
    assert well["sampling"] is True
    assert "samples" not in _gap_keys(well)


@needs_db
def test_passport_sets_equipment_and_protection(client, generated):
    document = _upload(client, generated / "documents" / "pasport_DEMO-011.pdf")
    assert document["status"] == "loaded", document["status_message"]
    well = _well(client, "DEMO-011")
    assert "покрытием" in well["equipment"]
    assert "защитное исполнение НКТ" in well["protection_means"]


@needs_db
def test_memo_is_not_recognised_and_loads_nothing(client, generated):
    document = _upload(client, generated / "documents" / "sluzhebnaya_zapiska.pdf")
    assert document["status"] == "not_recognized"
    assert document["status_label"] == "Не распознан"
    assert not document.get("loaded")


@needs_db
def test_invented_values_are_not_loaded(client, generated, fake):
    original = fake.respond

    def hallucinating(body):
        answer = original(body)
        return answer.replace("DEMO-002", "DEMO-001").replace("язвенная", "общая")

    fake.respond = hallucinating
    document = _upload(client, generated / "documents" / "akt_razbora_DEMO-002.pdf")
    assert document["status"] == "not_recognized"
    assert "base_act" in _gap_keys(_well(client, "DEMO-002"))


@needs_db
def test_inbox_folder(client, generated):
    from app.config import settings

    inbox = TEST_TMP / "inbox"
    target = inbox / "protokol.xlsx"
    shutil.copy(generated / "documents" / "protokol_khim_analiza_15.xlsx", target)
    old = time.time() - 60
    os.utime(target, (old, old))
    (inbox / "notes.txt").write_text("not a document")
    os.utime(inbox / "notes.txt", (old, old))
    summary = client.post("/api/v1/documents/inbox/scan").json()
    assert summary["accepted"] == 1 and summary["rejected"] == 1
    assert not target.exists()
    assert list((inbox / "processed").rglob("protokol.xlsx"))
    assert list((inbox / "failed").rglob("notes.txt.error.txt"))
    processed = client.post("/api/v1/documents/process-queue").json()["documents"]
    assert [item["status"] for item in processed] == ["loaded"]
    assert settings.docs_inbox_dir


@needs_db
def test_ocr_unavailable_is_retried_then_failed(client, generated, fake, monkeypatch):
    from app import db
    from app.config import settings
    from app.core import llm

    def broken(messages, **kwargs):
        if isinstance(messages[-1]["content"], list):
            raise llm.LLMError("LLM call failed (ConnectionError).")
        return fake.chat(messages, **kwargs)

    monkeypatch.setattr(llm, "chat", broken)
    document = _upload(client, generated / "documents" / "akt_demontazha_DEMO-005_scan.jpg")
    assert document["status"] == "queued" and "Повтор" in document["status_message"]
    for _ in range(settings.docs_max_attempts - 1):
        db.update_document(document["document_id"], started_at=db.datetime.now(db.timezone.utc) - timedelta(minutes=5))
        db.claim_next_document()
        from app.core.documents.pipeline import process_document
        row = process_document(document["document_id"])
    assert row["status"] == "failed"
    assert "OCR" in row["status_message"]


@needs_db
def test_ai_proposal_uses_rating_crews_and_reports_gaps(client, generated, fake):
    response = client.post("/api/v1/ingestion/crew-schedule?all_wells=true", json={"crew_availability": [
        {"crew_id": "КРС-7", "status": "available", "specialties": ["КРС"],
         "available_from": f"{TODAY.isoformat()}T08:00:00", "available_to": f"{(TODAY + timedelta(days=30)).isoformat()}T20:00:00"},
    ]})
    assert response.status_code == 201, response.text
    _upload(client, generated / "documents" / "sluzhebnaya_zapiska.pdf")
    proposal = client.post("/api/v1/rating/wells/DEMO-002/ai-proposal")
    assert proposal.status_code == 200, proposal.text
    body = proposal.json()
    well = _well(client, "DEMO-002")
    assert body["measure"] == well["measure_code"]
    assert body["crew"]["crew_id"] == "КРС-7"
    assert any("акта разбора" in item for item in body["missing_data"])
    assert well["ai_proposal"]["measure"] == body["measure"]
    assert "ai_proposal" not in client.get("/api/v1/rating").json()["wells"][0]

    # A crew that is not in the schedule is dropped, not invented.
    original = fake.respond
    fake.respond = lambda request: original(request).replace("КРС-7", "КРС-99")
    body = client.post("/api/v1/rating/wells/DEMO-002/ai-proposal").json()
    assert body["crew"] is None
    assert any("КРС-99" in note for note in body["checks"])


@needs_db
def test_scanned_pdf_and_source_table(client, generated, tmp_path):
    from PIL import Image

    pdf = tmp_path / "akt_demontazha_scan.pdf"
    Image.open(generated / "documents" / "akt_demontazha_DEMO-005_scan.jpg").convert("RGB").save(pdf, "PDF", resolution=150)
    document = _upload(client, pdf)
    assert document["status"] == "loaded", document["status_message"]
    assert document["text_method"] == "ocr"

    # A ХИМИК export dropped among documents goes through the usual source import.
    table = _upload(client, generated / "sample_khimik.xlsx")
    assert table["status"] == "loaded", table["status_message"]
    assert table["doc_type"] == "source:khimik"
    assert table["type_title"] == "Выгрузка KHIMIK"
    assert len(table["well_ids"]) >= 10


# --------------------------------------------------------------------------- mixed chemistry tables
def _chemistry_book(path, date_header="Дата", title=None, text_dates=True):
    from datetime import datetime as _dt
    from openpyxl import Workbook

    when = TODAY - timedelta(days=3)
    value = when.strftime("%d.%m.%Y") if text_dates else _dt.combine(when, _dt.min.time()).replace(hour=10)
    book = Workbook()
    sheet = book.active
    for line in title or []:
        sheet.append([line])
    if title:
        sheet.append([])
    sheet.append(["Скважина", date_header, "CO2, %", "pH", "Обводнённость, %", "Хлориды, мг/л", "Жёсткость, мг-экв/л", "Куст"])
    sheet.append(["DEMO-012", value, 1.4, 6.8, 47, 21000, 125, 7])
    sheet.append(["DEMO-013", value, 1.2, 7.1, 52, 19500, 98, 7])
    book.save(path)
    return path


def test_numbers_are_not_merged_across_table_cells():
    assert extract.value_agrees(21000.0, "number", "DEMO-012 | 16.10.2026 | 1.4 | 6.8 | 47 | 21000 | 125", "x")
    assert not extract.value_agrees(21000125.0, "number", "DEMO-012 | 47 | 21000 | 125", "x")
    assert extract.value_agrees(18500.0, "number", "Хлориды: 18 500 мг/л", "x")


def test_text_dates_in_excel():
    from app.core.excel_ingestion import _as_datetime

    assert _as_datetime("16.10.2026").date().isoformat() in {"2026-10-15", "2026-10-16"}   # local midnight in UTC
    assert _as_datetime("16.10.2026 10:30").hour == 7                                        # Moscow → UTC
    with pytest.raises(ValueError):
        _as_datetime("вчера")


def test_chemistry_only_table_is_a_khimik_sample(tmp_path):
    from app.core.documents.text import detect_profile
    from app.core.excel_ingestion import parse_excel

    path = _chemistry_book(tmp_path / "chem.xlsx")
    assert detect_profile(path.read_bytes(), path.name) == "khimik"
    parsed = parse_excel(path.read_bytes(), path.name, "khimik")
    assert parsed["rows_accepted"] == 2
    assert parsed["extra_columns"] == ["pH", "Жёсткость, мг-экв/л"]          # «Куст» is an identifier
    point = parsed["groups"]["DEMO-012"]["records"][0]
    assert point.extra_measurements == {"pH": 6.8, "Жёсткость, мг-экв/л": 125.0}


@needs_db
def test_mixed_chemistry_table_through_documents(client, tmp_path):
    document = _upload(client, _chemistry_book(tmp_path / "analizy.xlsx"))
    assert document["status"] == "loaded", document["status_message"]
    assert document["doc_type"] == "source:khimik"
    detail = client.get(f"/api/v1/documents/{document['document_id']}").json()
    assert detail["loaded"]["extra_columns"] == ["pH", "Жёсткость, мг-экв/л"]
    assert "Куст" in detail["loaded"]["unused_columns"]
    well = _well(client, "DEMO-012")
    assert well["sampling"] is True and "samples" not in _gap_keys(well)
    extras = {item["name"]: item for item in well["extra_measurements"]}
    assert extras["pH"]["value"] == 6.8 and extras["pH"]["source"] == "ХИМИК"
    assert extras["pH"]["date"][:10] == (TODAY - timedelta(days=3)).isoformat()


@needs_db
def test_protocol_with_other_indicators(client, tmp_path, fake):
    path = _chemistry_book(tmp_path / "protokol.xlsx", title=["ПРОТОКОЛ ХИМИЧЕСКОГО АНАЛИЗА № 21"])
    document = _upload(client, path)
    assert document["status"] == "loaded", document["status_message"]
    assert document["doc_type"] == "source:khimik"
    assert not [call for call in fake.calls if call["kind"] in {"extract", "classify"}]
    detail = client.get(f"/api/v1/documents/{document['document_id']}").json()
    assert detail["loaded"]["extra_columns"] == ["pH", "Жёсткость, мг-экв/л"]
    extras = {item["name"]: item for item in _well(client, "DEMO-012")["extra_measurements"]}
    assert extras["pH"]["value"] == 6.8 and extras["pH"]["source"] == "ХИМИК"
    # The AI agent sees them as context.
    from app.core.proposal_agent import build_context
    assert any(item["name"] == "pH" for item in build_context("DEMO-012")["other_measurements"])


# --------------------------------------------------------------------------- Excel without AI
def _book(path, rows):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Анализы"
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _custom_table(path, date_in_title=True, wells=("DEMO-012", "DEMO-013")):
    when = (TODAY - timedelta(days=2)).strftime("%d.%m.%Y")
    rows = [["Результаты анализа попутной воды"]]
    if date_in_title:
        rows.append([f"Дата отбора проб: {when}"])
    rows += [[], ["№ п/п", "№ скв.", "Обводн., %", "Cl-, мг/л", "Содержание CO2, %", "pH", "Fe общ., мг/л", "Эффективность ингибитора, %"]]
    for index, well in enumerate(wells, start=1):
        rows.append([index, well, 40 + index, f"{20 + index} 500", 1.1 + index / 10, 6.5 + index / 10, 0.3, 86])
    return _book(path, rows)


def test_excel_table_with_custom_headers_and_date_in_title(tmp_path):
    from app.core.documents.excel_reader import read_workbook

    result = read_workbook(_custom_table(tmp_path / "t.xlsx").read_bytes(), ["DEMO-012", "DEMO-013"])
    assert result["layout"] == "table" and result["accepted"] == 2
    record = result["records"][0]
    fields = {name: entry["value"] for name, entry in record["fields"].items() if entry["status"] == "ok"}
    assert record["well_id"] == "DEMO-012"
    assert fields["sample_date"] == (TODAY - timedelta(days=2)).isoformat()
    assert fields["water_cut_pct"] == 41 and fields["chlorides_mg_l"] == 21500 and fields["co2_pct"] == 1.2
    assert fields["inhibitor_efficiency"] == 0.86                     # percent -> fraction
    assert {item["name"] for item in fields["extra_measurements"]} == {"pH", "Fe общ., мг/л"}   # «№ п/п» skipped
    assert record["fields"]["chlorides_mg_l"]["quote"].startswith("Лист «Анализы», D5")


def test_excel_numeric_well_unknown_well_and_missing_date(tmp_path):
    from app.core.documents.excel_reader import not_found_reason, read_workbook

    path = _book(tmp_path / "n.xlsx", [["Скважина", "Дата", "Хлориды, мг/л"], [1234, TODAY, 21000], [5555, TODAY, 20000]])
    result = read_workbook(path.read_bytes(), ["1234"])
    assert result["accepted"] == 1 and result["records"][0]["well_id"] == "1234"
    assert "Скважины «5555» нет в справочнике." in result["records"][1]["issues"]
    no_date = read_workbook(_custom_table(tmp_path / "d.xlsx", date_in_title=False).read_bytes(), ["DEMO-012", "DEMO-013"])
    assert no_date["accepted"] == 0 and "Не найдена дата" in not_found_reason(no_date)
    nothing = read_workbook(_book(tmp_path / "x.xlsx", [["Итоги квартала"], ["План", 10]]).read_bytes(), ["DEMO-012"])
    assert "не найдена скважина" in not_found_reason(nothing)


def _form(path, well="DEMO-012"):
    when = (TODAY - timedelta(days=1)).strftime("%d.%m.%Y")
    return _book(path, [["ПРОТОКОЛ ИСПЫТАНИЙ № 7"], [f"Скважина № {well}"], ["Дата отбора", when], [],
                        ["Показатель", "Ед. изм.", "Значение"], ["Обводнённость", "%", 48], ["Хлориды", "мг/л", "21 000"],
                        ["pH", None, "6,9"], ["Сульфаты", "мг/л", 140], ["Номер пробы", None, 15]])


def test_excel_form_with_one_sample(tmp_path):
    from app.core.documents.excel_reader import read_workbook

    result = read_workbook(_form(tmp_path / "f.xlsx").read_bytes(), ["DEMO-012"])
    assert result["layout"] == "form" and result["accepted"] == 1
    fields = {name: entry["value"] for name, entry in result["records"][0]["fields"].items() if entry["status"] == "ok"}
    assert fields["water_cut_pct"] == 48 and fields["chlorides_mg_l"] == 21000
    assert {item["name"] for item in fields["extra_measurements"]} == {"pH", "Сульфаты, мг/л"}


@needs_db
def test_excel_documents_load_without_llm(client, tmp_path, monkeypatch):
    from app.config import settings
    from app.core import llm

    def no_llm(*args, **kwargs):
        raise AssertionError("Excel must not call the LLM")

    monkeypatch.setattr(llm, "chat", no_llm)
    monkeypatch.setattr(type(settings), "extraction_model", property(lambda self: ""))
    table = _upload(client, _custom_table(tmp_path / "analizy.xlsx"))
    assert table["status"] == "loaded", table["status_message"]
    assert table["doc_type"] == "excel:table" and table["type_title"] == "Таблица показателей (Excel)"
    well = _well(client, "DEMO-012")
    assert well["sampling"] is True and "samples" not in _gap_keys(well)
    assert {"pH", "Fe общ., мг/л"} <= {item["name"] for item in well["extra_measurements"]}
    form = _upload(client, _form(tmp_path / "forma.xlsx", well="DEMO-013"))
    assert form["status"] == "loaded" and form["doc_type"] == "excel:form", form["status_message"]
    unknown = _upload(client, _custom_table(tmp_path / "chuzhie.xlsx", wells=("X-1", "X-2")))
    assert unknown["status"] == "not_recognized" and "нет в справочнике" in unknown["status_message"]
    detail = client.get(f"/api/v1/documents/{table['document_id']}").json()
    assert detail["field_labels"]["chlorides_mg_l"] == "Хлориды, мг/л"


@needs_db
def test_reupload_after_failure_processes_again(client, generated):
    from app import db

    path = generated / "documents" / "protokol_khim_analiza_15.xlsx"
    first = _upload(client, path)
    assert first["status"] == "loaded"
    # As if it had failed before (e.g. uploaded before the Excel reader existed).
    db.update_document(first["document_id"], status="failed", status_message="Языковая модель не настроена")
    with open(path, "rb") as handle:
        body = client.post("/api/v1/documents?wait=true", files=[("files", ("protocol khim analyzs.xlsx", handle))]).json()
    again = body["documents"][0]
    assert again["document_id"] == first["document_id"]
    assert again["status"] == "loaded" and again["requeued"] is True
    assert body["requeued"] == 1 and body["duplicates"] == 0
    # A loaded file sent again is just a duplicate.
    with open(path, "rb") as handle:
        body = client.post("/api/v1/documents?wait=true", files=[("files", (path.name, handle))]).json()
    assert body["duplicates"] == 1 and body["requeued"] == 0


@needs_db
def test_clear_document_list(client, generated):
    memo = _upload(client, generated / "documents" / "sluzhebnaya_zapiska.pdf")
    protocol = _upload(client, generated / "documents" / "protokol_khim_analiza_15.xlsx")
    assert memo["status"] == "not_recognized" and protocol["status"] == "loaded"
    assert client.delete("/api/v1/documents").status_code == 422
    result = client.delete("/api/v1/documents?status=failed,not_recognized").json()
    assert result == {"deleted": 1, "wells": []}
    assert [row["document_id"] for row in client.get("/api/v1/documents").json()["documents"]] == [protocol["document_id"]]
    assert "samples" not in _gap_keys(_well(client, "DEMO-012"))
    result = client.delete("/api/v1/documents?all=true").json()
    assert result["deleted"] == 1 and result["wells"] == ["DEMO-012", "DEMO-013"]
    assert client.get("/api/v1/documents").json()["documents"] == []
    assert "samples" in _gap_keys(_well(client, "DEMO-012"))        # its data went with it


# --------------------------------------------------------------------------- Реестр коррозии НКТ
REGISTRY_HEAD = ["№", "ДО", "Месторождение", "Куст", "Скважина", "Глубина спуска насоса, м", "Пласт",
                 "Факты выявления коррозии НКТ на скважине", "ННО, сут", "Режим работы ЭЦН", "Признак новизны НКТ",
                 "Диаметр НКТ", "Группа прочности поврежденной НКТ",
                 "Состав подвески(состоит из 3 Материал НКТ группа прочности вверх, середина, низ)", None, None,
                 "Температура на устье, оС", "Температура пластовая, оС", "Давление на устье, атм",
                 "Средний дебит жидкости, м3", "Средняя обводненность, %", "Плотность воды, кг/м3", "Ca2+", "Mg2+",
                 "Na+ K+", "Feобщ", "Cl-", "SO42-", "HCO3-", "CO32", "Общая минерализация, мг/л", "pH воды, ед",
                 "Среднее значение КВЧ, мг/л", "CO2, мл/г", "H2S, мг/л", "CO2, % мол.", "H2S % мол.",
                 "Газовый фактор м3/т", "марка стали", "Завод", "Отказы за весь период, кол-во", "Тип повреждения",
                 "Глубина от устья, где произошел отказ, м", "Толщина стенки поврежденной НКТ, мм"]


def _registry(path, rows, title=None):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Реестр"
    if title:
        sheet.append([title])
    sheet.append(REGISTRY_HEAD)
    sub = [None] * len(REGISTRY_HEAD)
    sub[13:16] = ["верх", "середина", "низ"]
    sheet.append(sub)
    sheet.append(list(range(1, len(REGISTRY_HEAD) + 1)))      # column numbers under the header
    header_row = 2 if title else 1
    sheet.merge_cells(start_row=header_row, start_column=14, end_row=header_row, end_column=16)
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _registry_row(no, well, field="Южное", water_cut=92, chlorides=9800):
    return [no, "ООО «Нефть-Юг»", field, "12", well, 2150, "ЮВ1", "да", 340, "постоянный", "новые", 73, "К",
            "К", "Е", "Н80", 38, 92, 14.5, 85, water_cut, 1012, 420, 95, 6100, 12.4, chlorides, 45, 610, 0, 17500,
            6.8, 120, 0.8, 2.1, 1.9, 0.01, 85, "Ст20", "ТМК", 3, "свищ", 1840, 5.5]


def test_registry_columns_and_two_row_header(tmp_path):
    from app.core.documents import nkt_registry

    when = (TODAY - timedelta(days=3)).strftime("%d.%m.%Y")
    path = _registry(tmp_path / "r.xlsx", [_registry_row(1, "DEMO-001"), _registry_row(2, "7781"),
                                           _registry_row(3, "7781", field="Северное"),
                                           [4, "ООО «Нефть-Юг»", "Южное", "14", "9001"]],
                     title=f"Реестр фактов коррозии НКТ по состоянию на {when}")
    content = path.read_bytes()
    assert nkt_registry.detect(content)
    assert not nkt_registry.detect(_custom_table(tmp_path / "t.xlsx").read_bytes())
    result = nkt_registry.read_workbook(content, ["DEMO-001"], TODAY)
    assert result["date"] == TODAY - timedelta(days=3) and result["date_note"] is None
    [record] = result["records"]
    values = {key: field["value"] for key, field in record["fields"].items()}
    assert record["well_id"] == "DEMO-001" and not record["new_well"]
    assert values["string_top"] == "К" and values["string_middle"] == "Е" and values["string_bottom"] == "Н80"
    assert values["cl"] == 9800 and values["co2_mol_pct"] == 1.9 and values["co2_ml_g"] == 0.8
    assert values["h2s_mg_l"] == 2.1 and values["h2s_mol_pct"] == 0.01 and values["co3"] == 0
    assert values["nno_days"] == 340 and values["failures_count"] == 3 and values["failure_depth_m"] == 1840
    assert values["corrosion_facts"] == "да" and values["damage_type"] == "свищ" and values["steel_grade"] == "Ст20"
    assert "row_no" not in values
    assert record["fields"]["cl"]["quote"].startswith("Лист «Реестр», AA")
    assert result["columns"]["system"] == ["Месторождение", "Средний дебит жидкости, м3", "Средняя обводненность, %",
                                           "Cl-", "CO2, % мол."]
    issues = {issue for item in result["rejected"] for issue in item["issues"]}
    assert any("разных месторождениях" in issue for issue in issues)            # 7781 on two fields
    assert "В строке нет данных, кроме расположения скважины." in issues          # 9001: location only
    no_date = nkt_registry.read_workbook(_registry(tmp_path / "n.xlsx", [_registry_row(1, "5")]).read_bytes(), [], TODAY)
    assert no_date["date"] == TODAY and "дата" in no_date["date_note"].casefold()
    assert no_date["records"][0]["new_well"] is False                          # no well list: taken as is


@needs_db
def test_registry_through_documents(client, tmp_path, monkeypatch):
    from app import db
    from app.config import settings
    from app.core import llm

    def no_llm(*args, **kwargs):
        raise AssertionError("The registry must not call the LLM")

    monkeypatch.setattr(llm, "chat", no_llm)
    monkeypatch.setattr(type(settings), "extraction_model", property(lambda self: ""))
    path = _registry(tmp_path / "reestr_nkt.xlsx", [_registry_row(1, "DEMO-012", water_cut=93, chlorides=31000),
                                                    _registry_row(2, "NEW-77")])
    document = _upload(client, path)
    assert document["status"] == "loaded", document["status_message"]
    assert document["doc_type"] == "excel:nkt_registry" and document["type_title"] == "Реестр коррозии НКТ (Excel)"
    assert document["well_ids"] == ["DEMO-012", "NEW-77"] and document["doc_date"] == TODAY.isoformat()
    assert "Новых скважин в системе: 1" in document["status_message"]
    assert "extraction" not in client.get("/api/v1/documents").json()["documents"][0]   # the list stays light
    detail = client.get(f"/api/v1/documents/{document['document_id']}").json()
    assert detail["extraction"]["records_accepted"] == 2 and detail["field_labels"]["string_top"] == "Состав подвески: верх"

    well = _well(client, "DEMO-012")
    registry = dict(well["nkt_registry"]["rows"][0])
    assert registry["Тип повреждения"] == "свищ" and registry["Состав подвески: низ"] == "Н80"
    point = [item for item in db.load_telemetry("DEMO-012") if getattr(item, "source_id", None) == "DOC_NKT"]
    assert len(point) == 1 and point[0].chlorides_mg_l == 31000 and point[0].water_cut_pct == 93
    assert point[0].co2_pct == 1.9 and point[0].liquid_rate_m3_day == 85
    assert "Реестр коррозии НКТ (документ)" in well["provenance"]["corrosive_environment"]["source"]
    assert _well(client, "NEW-77")["field"] == "Южное"                          # a new well gets a rating row

    assert client.delete(f"/api/v1/documents/{document['document_id']}").status_code == 200
    assert _well(client, "DEMO-012")["nkt_registry"] is None
    assert not [item for item in db.load_telemetry("DEMO-012") if getattr(item, "source_id", None) == "DOC_NKT"]
