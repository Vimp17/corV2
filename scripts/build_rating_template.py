"""Build the well-rating Excel template (approved layout) with synthetic sample data.

    python scripts/build_rating_template.py data/sample_rating.xlsx

The «Рейтинг» sheet is also the upload format of the "rating" source profile: the dashboard
and POST /api/v1/ingestion/excel read its input columns and recompute the index, measure and
deadlines themselves. Measure names and thresholds come from app.core.rating and the settings,
so the template always matches the backend. Recalculate formulas by opening the file in Excel
or LibreOffice (openpyxl does not compute them).
"""
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.config import settings  # noqa: E402
from app.core import rating  # noqa: E402

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Synthetic sample: well, field, equipment, status, sampling, environment, tonnage, НОО, demount act,
# corrosion sign, base act, base result, category, НТЦ, math, plan date, Промайн document, comment.
SAMPLE_WELLS = [
    ["2215Б", "Озёрное", "УЭЦН-5-80, НКТ-73", "Не в работе", "Да", "Коррозионная", 3.0, 38, "Есть", "Есть", "Есть", "Коррозия есть", "Коррозионная", 0.93, 0.88, "2026-10-21", "Сформирован", "Язвенная коррозия НКТ по акту разбора"],
    ["1131", "Северное", "УЭЦН-5А-125, НКТ-73", "В работе", "Да", "Коррозионная", 2.5, 52, "Есть", "Есть", "Нет", "—", "Коррозионная", 0.85, 0.88, "2026-10-23", "Не сформирован", ""],
    ["3307", "Восточное", "УЭЦН-5-45, НКТ-60", "В работе", "Да", "Коррозионная", 2.0, 61, "Нет", "—", "Нет", "—", "Коррозионная", 0.85, 0.62, "2026-10-28", "Не сформирован", "Модели расходятся: проверить пробы"],
    ["1046", "Северное", "УШГН-57, НКТ-73", "Не в работе", "Нет", "Коррозионная", 1.5, 74, "Есть", "Есть", "Есть", "Коррозия есть", "Коррозионная", 0.79, 0.81, "2026-11-05", "Направлен в ЭРА", "Ждёт бригаду КРС"],
    ["2241", "Озёрное", "УЭЦН-5-80, НКТ-73", "В работе", "Да", "Коррозионная", 1.5, 88, "Есть", "Нет", "Есть", "Коррозия есть", "Среднекоррозионная", 0.76, 0.71, "2026-11-09", "—", ""],
    ["3312", "Восточное", "УЭЦН-5-60, НКТ-73", "В работе", "Да", "Коррозионная", 1.0, 95, "Нет", "—", "Нет", "—", "Среднекоррозионная", 0.66, 0.69, "2026-11-16", "—", ""],
    ["1102", "Северное", "УЭЦН-5-45, НКТ-60", "В работе", "Нет", "Некоррозионная", 0.0, 140, "Есть", "Есть", "Нет", "—", "Среднекоррозионная", 0.64, 0.52, "2026-11-12", "—", "Нет свежих проб"],
    ["2208", "Озёрное", "УШГН-44, НКТ-60", "В работе", "Да", "Коррозионная", 1.0, 102, "Нет", "—", "Нет", "—", "Среднекоррозионная", 0.58, 0.55, "2026-11-02", "—", ""],
    ["3320", "Восточное", "УЭЦН-5-80, НКТ-73", "В работе", "Да", "Коррозионная", 1.0, 115, "Есть", "Нет", "Есть", "Коррозии нет", "Среднекоррозионная", 0.49, 0.55, "2026-11-11", "—", ""],
    ["1158", "Северное", "УЭЦН-5А-125, НКТ-73", "В работе", "Да", "Некоррозионная", 0.5, 130, "Нет", "—", "Нет", "—", "Среднекоррозионная", 0.51, 0.47, "2026-11-18", "—", ""],
    ["2219", "Озёрное", "УЭЦН-5-60, НКТ-73", "В работе", "Нет", "Коррозионная", 0.5, 160, "Нет", "—", "Нет", "—", "Среднекоррозионная", 0.47, 0.4, "2026-11-25", "—", ""],
    ["3301", "Восточное", "УШГН-57, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 190, "Есть", "Нет", "Есть", "Коррозии нет", "Некоррозионная", 0.38, 0.44, "", "—", ""],
    ["1077", "Северное", "УЭЦН-5-45, НКТ-60", "В работе", "Да", "Некоррозионная", 0.0, 210, "Нет", "—", "Нет", "—", "Некоррозионная", 0.41, 0.36, "2026-11-26", "—", ""],
    ["2230", "Озёрное", "УЭЦН-5-80, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 240, "Нет", "—", "Нет", "—", "Некоррозионная", 0.37, 0.33, "", "—", ""],
    ["3315", "Восточное", "УЭЦН-5-60, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 255, "Нет", "—", "Нет", "—", "Некоррозионная", 0.31, 0.35, "", "—", ""],
    ["1119", "Северное", "УШГН-44, НКТ-60", "В работе", "Нет", "Некоррозионная", 0.0, 270, "Нет", "—", "Нет", "—", "Некоррозионная", 0.33, 0.29, "", "—", ""],
    ["2244", "Озёрное", "УЭЦН-5-45, НКТ-60", "Не в работе", "Нет", "Некоррозионная", 0.0, 300, "Есть", "Нет", "Есть", "Коррозии нет", "Некоррозионная", 0.28, 0.3, "", "—", "Остановлена по геологии"],
    ["3326", "Восточное", "УЭЦН-5А-125, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 310, "Нет", "—", "Нет", "—", "Некоррозионная", 0.26, 0.22, "", "—", ""],
    ["1085", "Северное", "УЭЦН-5-80, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 330, "Нет", "—", "Нет", "—", "Некоррозионная", 0.21, 0.24, "", "—", ""],
    ["2202", "Озёрное", "УШГН-57, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 350, "Нет", "—", "Нет", "—", "Некоррозионная", 0.19, 0.17, "", "—", ""],
    ["3309", "Восточное", "УЭЦН-5-45, НКТ-60", "В работе", "Нет", "Некоррозионная", 0.0, 365, "Нет", "—", "Нет", "—", "Некоррозионная", 0.15, 0.18, "", "—", ""],
    ["1140", "Северное", "УЭЦН-5-60, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 380, "Нет", "—", "Нет", "—", "Некоррозионная", 0.12, 0.14, "", "—", ""],
    ["2236", "Озёрное", "УЭЦН-5-80, НКТ-73", "В работе", "Да", "Некоррозионная", 0.0, 400, "Нет", "—", "Нет", "—", "Некоррозионная", 0.09, 0.11, "", "—", ""],
    ["3330", "Восточное", "УШГН-44, НКТ-60", "В работе", "Да", "Некоррозионная", 0.0, 420, "Нет", "—", "Нет", "—", "Некоррозионная", 0.07, 0.06, "", "—", ""]
]

data = {"calc_date": "2026-10-14", "gantt_start": "2026-10-12", "wells": SAMPLE_WELLS}
out = sys.argv[1] if len(sys.argv) > 1 else "data/sample_rating.xlsx"
wells = data["wells"]
LBL = rating.MEASURE_LABELS
CODE = rating.MEASURE_CODES
N = len(wells)
FIRST = 5
LAST = FIRST + N - 1

FONT = "Arial"
NAVY, NAVY2, BLUE = "0E2A47", "16375C", "0072CE"
f_base = Font(name=FONT, size=10)
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=14, bold=True, color=NAVY)
f_note = Font(name=FONT, size=9, italic=True, color="5A6B7C")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_group = PatternFill("solid", fgColor=NAVY2)
fill_input = PatternFill("solid", fgColor="FFF9DB")
fill_calc = PatternFill("solid", fgColor="EEF3F9")
thin = Side(style="thin", color="DDE3EA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Semantic fills shared by the rating and ППР sheets (match the dashboard palette).
SEM = {
    "ПКР": ("FDECEA", "B71C1C"),  # planned repair package + ERA repair
    "РИ": ("FDEFE3", "BF4300"),   # engineer decision
    "УД": ("FFF4D6", "8A5300"),   # shock dosing
    "НЕТ": ("E8F5E9", "2E7D32"),  # no measures required
}

wb = Workbook()

# --------------------------------------------------------------------------- Параметры
P = wb.active
P.title = "Параметры"
P["A1"] = "Параметры рейтинга и графика ППР"
P["A1"].font = f_title
P["A2"] = ("Синие значения на жёлтом фоне можно менять: рейтинг, мероприятия, сроки и график "
           "пересчитываются автоматически.")
P["A2"].font = f_note
rows = [
    ("Дата расчёта рейтинга", date.fromisoformat(data["calc_date"]), "DD.MM.YYYY",
     "От неё считаются крайние сроки мероприятий."),
    ("Объединение моделей НТЦ и мат.", "Максимум", None,
     "Максимум — консервативно (берётся более тревожная оценка); Среднее — компромисс."),
    ("Начало графика ППР (понедельник)", date.fromisoformat(data["gantt_start"]), "DD.MM.YYYY",
     "Первая неделя листа «График ППР»; 12 недель вперёд."),
]
for i, (label, value, fmt, note) in enumerate(rows, start=4):
    P.cell(i, 1, label).font = f_base
    c = P.cell(i, 3, value)
    c.font, c.fill, c.border = f_input, fill_input, border
    if fmt:
        c.number_format = fmt
    P.cell(i, 4, note).font = f_note
dv_method = DataValidation(type="list", formula1='"Максимум,Среднее"', allow_blank=False)
P.add_data_validation(dv_method)
dv_method.add("C5")

P["A8"] = "Правила мероприятий по индексу приоритета (0–100)"
P["A8"].font = f_bold
heads = ["Мероприятие", "Код", "Индекс от", "Индекс до", "Срок, дней", "Основание"]
for j, h in enumerate(heads, start=1):
    c = P.cell(9, j, h)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
rules = [
    (LBL["repair"], CODE["repair"], settings.rating_repair_from, "=100", settings.rating_days_repair,
     "Порог 80+ — из требования: плановый комплекс мероприятий по ремонту, скважина включается "
     "в ЭРА ремонта, Промайн формирует документ. Срок в днях — предложение, уточнить."),
    (LBL["engineer"], CODE["engineer"], "=D12+1", "=C10-1", settings.rating_days_engineer,
     "Мероприятие для диапазона 61–79 в требовании не задано — уточнить. Срок — предложение."),
    (LBL["ud"], CODE["ud"], settings.rating_ud_from, settings.rating_ud_to, settings.rating_days_ud,
     "Диапазон 40–60 — из требования. Расшифровка «УД» как ударной дозировки — допущение. "
     "Срок — предложение."),
    (LBL["none"], CODE["none"], 0, "=C12-1", None, "Ниже порога УД мероприятия не требуются."),
]
for i, (name, code, lo, hi, days, why) in enumerate(rules, start=10):
    P.cell(i, 1, name)
    P.cell(i, 2, code)
    P.cell(i, 3, lo)
    P.cell(i, 4, hi)
    P.cell(i, 5, days)
    P.cell(i, 6, why)
    for j in range(1, 7):
        c = P.cell(i, j)
        c.border, c.alignment = border, (left if j in (1, 6) else center)
        c.font = f_base
    for j, v in ((3, lo), (4, hi), (5, days)):
        if not (isinstance(v, str) and v.startswith("=")) and v is not None:
            P.cell(i, j).font, P.cell(i, j).fill = f_input, fill_input
P["A15"] = ("Если индекс выше верхней границы УД, но ниже порога планового комплекса по ремонту, "
            "назначается «Решение инженера». Нижние границы считаются формулами от соседних порогов. "
            "Пороги и сроки в системе задаются переменными MAI_RATING_* — держите их одинаковыми.")
P["A15"].font = f_note
for col, w in zip("ABCDEF", (36, 8, 13, 11, 11, 78)):
    P.column_dimensions[col].width = w
P.row_dimensions[10].height = 42
P.row_dimensions[12].height = 30

CALC, METHOD, GSTART = "Параметры!$C$4", "Параметры!$C$5", "Параметры!$C$6"
TO_FROM, UD_FROM, UD_TO = "Параметры!$C$10", "Параметры!$C$12", "Параметры!$D$12"
D_TO, D_RI, D_UD = "Параметры!$E$10", "Параметры!$E$11", "Параметры!$E$12"
L_TO, L_RI, L_UD, L_MON = ("Параметры!$A$10", "Параметры!$A$11", "Параметры!$A$12",
                           "Параметры!$A$13")
C_TO, C_RI, C_UD = "Параметры!$B$10", "Параметры!$B$11", "Параметры!$B$12"

# --------------------------------------------------------------------------- Рейтинг
R = wb.create_sheet("Рейтинг", 0)
R["A1"] = "Рейтинг скважин по риску коррозии"
R["A1"].font = f_title
R["A2"] = ("Синий текст на жёлтом — ввод (из источников или вручную), чёрный на голубом — формулы. "
           "Пример заполнен синтетическими данными. Строки отсортированы по индексу.")
R["A2"].font = f_note

cols = [
    # key, header, width, kind, group
    ("rank", "Место", 7, "calc", "Скважина"),
    ("well", "№ скважины", 11, "in", "Скважина"),
    ("field", "Месторождение", 16, "in", "Скважина"),
    ("equip", "Спущенное оборудование", 22, "in", "Скважина"),
    ("status", "Статус", 12, "in", "Скважина"),
    ("sampling", "Отбор проб", 10, "in", "Среда и обработка"),
    ("env", "Среда", 15, "in", "Среда и обработка"),
    ("tonnage", "Тоннаж, т", 10, "in", "Среда и обработка"),
    ("noo", "НОО, сут", 10, "in", "Среда и обработка"),
    ("act_dem", "Акт демонтажа", 11, "in", "Демонтаж и разбор"),
    ("corr_sign", "Признак коррозии", 11, "in", "Демонтаж и разбор"),
    ("act_base", "Акт разбора на базе", 11, "in", "Демонтаж и разбор"),
    ("base_res", "Результат разбора", 14, "in", "Демонтаж и разбор"),
    ("category", "Категория скважины", 21, "in", "Оценка"),
    ("ntc", "Модель НТЦ (0–1)", 10, "in", "Оценка"),
    ("math", "Мат. модель (0–1)", 10, "in", "Оценка"),
    ("index", "Индекс приоритета (0–100)", 11, "calc", "Оценка"),
    ("measure", "Мероприятие", 26, "calc", "Решение и график ППР"),
    ("era", "ЭРА ремонта", 11, "calc", "Решение и график ППР"),
    ("promine", "Документ Промайн", 16, "in", "Решение и график ППР"),
    ("deadline", "Крайний срок", 12, "calc", "Решение и график ППР"),
    ("plan", "План ППР (ДДО)", 12, "in", "Решение и график ППР"),
    ("plan_status", "Статус плана", 15, "calc", "Решение и график ППР"),
    ("comment", "Комментарий", 34, "in", ""),
]
col = {key: get_column_letter(i) for i, (key, *_rest) in enumerate(cols, start=1)}

# group header row 3
groups = []
for i, (_k, _h, _w, _kind, g) in enumerate(cols, start=1):
    if groups and groups[-1][0] == g:
        groups[-1][2] = i
    else:
        groups.append([g, i, i])
for g, a, b in groups:
    c = R.cell(3, a, g)
    c.font, c.fill, c.alignment = f_head, fill_group, center
    for j in range(a, b + 1):
        R.cell(3, j).fill = fill_group
        R.cell(3, j).border = border
    if b > a:
        R.merge_cells(start_row=3, start_column=a, end_row=3, end_column=b)
for i, (_k, h, w, _kind, _g) in enumerate(cols, start=1):
    c = R.cell(4, i, h)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
    R.column_dimensions[get_column_letter(i)].width = w
R.row_dimensions[4].height = 44

notes = {
    "tonnage": "Тоннаж (поле из требования). Единицы (т) — допущение, уточнить.",
    "noo": "НОО — поле из требования; расшифровку и единицы (здесь сутки) уточнить.",
    "corr_sign": "Признак коррозии при демонтаже. «—», если акта демонтажа нет.",
    "base_res": "Заполняется, только если есть акт разбора на базе.",
    "index": "=ОКРУГЛ(Максимум или Среднее(НТЦ; Мат.) × 100). Способ — на листе «Параметры».",
    "measure": "По порогам с листа «Параметры»: 80+ плановый комплекс по ремонту и ЭРА, 61–79 "
               "решение инженера, 40–60 УД, ниже 40 мероприятия не требуются. При загрузке в систему "
               "пересчитывается.",
    "promine": "Только для скважин 80+: Промайн формирует документ на включение в ЭРА ремонта.",
    "deadline": "Дата расчёта + срок мероприятия с листа «Параметры».",
    "plan": "Плановая дата мероприятия по графику ППР (ДДО). Ввод.",
}
for key, text in notes.items():
    R[f"{col[key]}4"].comment = Comment(text, "Шаблон")

keys = [k for k, *_ in cols]
in_map = {"well": 0, "field": 1, "equip": 2, "status": 3, "sampling": 4, "env": 5, "tonnage": 6,
          "noo": 7, "act_dem": 8, "corr_sign": 9, "act_base": 10, "base_res": 11, "category": 12,
          "ntc": 13, "math": 14, "plan": 15, "promine": 16, "comment": 17}
for n, w in enumerate(wells):
    r = FIRST + n
    for key, idx in in_map.items():
        v = w[idx]
        if key == "plan":
            v = date.fromisoformat(v) if v else None
        R[f"{col[key]}{r}"] = v
    q, o, p = col["index"], col["ntc"], col["math"]
    R[f"{col['rank']}{r}"] = f'=IF({q}{r}="","",RANK({q}{r},${q}${FIRST}:${q}${LAST},0))'
    R[f"{q}{r}"] = (f'=IF(COUNT({o}{r}:{p}{r})=0,"",ROUND(IF({METHOD}="Среднее",'
                    f'AVERAGE({o}{r}:{p}{r}),MAX({o}{r}:{p}{r}))*100,0))')
    R[f"{col['measure']}{r}"] = (f'=IF({q}{r}="","",IF({q}{r}>={TO_FROM},{L_TO},IF({q}{r}>{UD_TO},'
                                 f'{L_RI},IF({q}{r}>={UD_FROM},{L_UD},{L_MON}))))')
    R[f"{col['era']}{r}"] = f'=IF({q}{r}="","",IF({q}{r}>={TO_FROM},"Включить","—"))'
    R[f"{col['deadline']}{r}"] = (f'=IF({q}{r}="","",IF({q}{r}>={TO_FROM},{CALC}+{D_TO},'
                                  f'IF({q}{r}>{UD_TO},{CALC}+{D_RI},IF({q}{r}>={UD_FROM},'
                                  f'{CALC}+{D_UD},""))))')
    v_, u_ = col["plan"], col["deadline"]
    R[f"{col['plan_status']}{r}"] = (f'=IF(OR({q}{r}="",{q}{r}<{UD_FROM}),"—",IF({v_}{r}="",'
                                     f'"Не запланировано",IF({v_}{r}>{u_}{r},"Позже срока","В срок")))')
    for i, (key, _h, _w, kind, _g) in enumerate(cols, start=1):
        c = R.cell(r, i)
        c.border = border
        if kind == "in":
            c.font, c.fill = f_input, fill_input
        else:
            c.font, c.fill = f_base, fill_calc
        c.alignment = left if key in ("equip", "measure", "comment", "category", "field") else center
    R[f"{col['ntc']}{r}"].number_format = "0.00"
    R[f"{col['math']}{r}"].number_format = "0.00"
    R[f"{col['tonnage']}{r}"].number_format = "0.0"
    R[f"{col['noo']}{r}"].number_format = "0"
    R[f"{col['index']}{r}"].number_format = "0"
    R[f"{col['deadline']}{r}"].number_format = "DD.MM.YYYY"
    R[f"{col['plan']}{r}"].number_format = "DD.MM.YYYY"
    R[f"{col['index']}{r}"].font = Font(name=FONT, size=10, bold=True)

R[f"A{LAST + 2}"] = (f"Строк в примере: {N}. Новые скважины добавляйте ниже, копируя последнюю строку "
                     "(формулы в голубых ячейках протянутся вместе с ней) и расширяя диапазон рейтинга "
                     "в колонке «Место».")
R[f"A{LAST + 2}"].font = f_note

# Dropdowns on input columns
lists = {
    "status": "В работе,Не в работе",
    "sampling": "Да,Нет",
    "env": "Коррозионная,Некоррозионная",
    "act_dem": "Есть,Нет",
    "corr_sign": "Есть,Нет,—",
    "act_base": "Есть,Нет",
    "base_res": "Коррозия есть,Коррозии нет,—",
    "category": "Некоррозионная,Среднекоррозионная,Коррозионная",
    "promine": "Не сформирован,Сформирован,Направлен в ЭРА,—",
}
for key, values in lists.items():
    dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
    R.add_data_validation(dv)
    dv.add(f"{col[key]}{FIRST}:{col[key]}{FIRST + 500}")
for key in ("ntc", "math"):
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                        allow_blank=True, error="Оценка модели должна быть от 0 до 1.",
                        showErrorMessage=True)
    R.add_data_validation(dv)
    dv.add(f"{col[key]}{FIRST}:{col[key]}{FIRST + 500}")

# Conditional formatting: measure / index / plan status / evidence
rng = lambda key: f"{col[key]}{FIRST}:{col[key]}{LAST}"  # noqa: E731
q = col["index"]
for key in ("index", "measure", "era"):
    for cond, sem in ((f"${q}{FIRST}>={TO_FROM}", "ПКР"),
                      (f"AND(${q}{FIRST}>{UD_TO},${q}{FIRST}<{TO_FROM})", "РИ"),
                      (f"AND(${q}{FIRST}>={UD_FROM},${q}{FIRST}<={UD_TO})", "УД"),
                      (f"AND(ISNUMBER(${q}{FIRST}),${q}{FIRST}<{UD_FROM})", "НЕТ")):
        bg, fg = SEM[sem]
        R.conditional_formatting.add(rng(key), FormulaRule(
            formula=[cond], fill=PatternFill("solid", fgColor=bg),
            font=Font(name=FONT, color=fg, bold=(key != "era"))))
w_ = col["plan_status"]
R.conditional_formatting.add(rng("plan_status"), FormulaRule(
    formula=[f'{w_}{FIRST}="Позже срока"'], fill=PatternFill("solid", fgColor="FDECEA"),
    font=Font(name=FONT, color="B71C1C", bold=True)))
R.conditional_formatting.add(rng("plan_status"), FormulaRule(
    formula=[f'{w_}{FIRST}="Не запланировано"'], fill=PatternFill("solid", fgColor="FFF4D6"),
    font=Font(name=FONT, color="8A5300", bold=True)))
for key, bad in (("env", "Коррозионная"), ("corr_sign", "Есть"), ("base_res", "Коррозия есть"),
                 ("category", "Коррозионная")):
    c_ = col[key]
    R.conditional_formatting.add(rng(key), FormulaRule(
        formula=[f'{c_}{FIRST}="{bad}"'], font=Font(name=FONT, color="B71C1C", bold=True)))

R.freeze_panes = f"C{FIRST}"
R.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{LAST}"
R.sheet_view.zoomScale = 90

# --------------------------------------------------------------------------- График ППР
G = wb.create_sheet("График ППР", 1)
G["A1"] = "График ППР (ДДО) — мероприятия по рейтингу на 12 недель"
G["A1"].font = f_title
G["A2"] = ("Код мероприятия стоит в неделе плановой даты; светлая заливка — окно от даты расчёта до "
           "крайнего срока. «!» — план позже срока. Неделя с датой расчёта выделена в заголовке.")
G["A2"].font = f_note
gheads = ["Место", "№ скважины", "Месторождение", "Индекс", "Код", "Крайний срок", "План (ДДО)"]
WEEKS = 12
for j, h in enumerate(gheads, start=1):
    c = G.cell(4, j, h)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
    G.cell(3, j).fill = fill_head
for k in range(WEEKS):
    j = len(gheads) + 1 + k
    L = get_column_letter(j)
    top = G.cell(3, j, f"Нед. {k + 1}")
    top.font, top.fill, top.alignment, top.border = f_head, fill_group, center, border
    c = G.cell(4, j, f"={GSTART}+{7 * k}")
    c.number_format = "DD.MM"
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
    G.column_dimensions[L].width = 8.5
G.row_dimensions[4].height = 30
for j, w in enumerate((7, 11, 16, 8, 7, 12, 12), start=1):
    G.column_dimensions[get_column_letter(j)].width = w

first_week = get_column_letter(len(gheads) + 1)
last_week = get_column_letter(len(gheads) + WEEKS)
for n in range(N):
    r, src = FIRST + n, FIRST + n
    rq = f"Рейтинг!{col['index']}{src}"
    G[f"A{r}"] = f"=Рейтинг!{col['rank']}{src}"
    G[f"B{r}"] = f"=Рейтинг!{col['well']}{src}"
    G[f"C{r}"] = f"=Рейтинг!{col['field']}{src}"
    G[f"D{r}"] = f"={rq}"
    G[f"E{r}"] = (f'=IF({rq}="","",IF({rq}>={TO_FROM},{C_TO},IF({rq}>{UD_TO},{C_RI},'
                  f'IF({rq}>={UD_FROM},{C_UD},""))))')
    G[f"F{r}"] = f"=Рейтинг!{col['deadline']}{src}"
    G[f"G{r}"] = f'=IF(Рейтинг!{col["plan"]}{src}="","",Рейтинг!{col["plan"]}{src})'
    for j in range(1, len(gheads) + 1):
        c = G.cell(r, j)
        c.font, c.border, c.alignment = f_link, border, (left if j == 3 else center)
    G[f"F{r}"].number_format = G[f"G{r}"].number_format = "DD.MM.YYYY"
    for k in range(WEEKS):
        L = get_column_letter(len(gheads) + 1 + k)
        G[f"{L}{r}"] = (f'=IF(AND(ISNUMBER($G{r}),$G{r}>={L}$4,$G{r}<{L}$4+7),'
                        f'$E{r}&IF(AND(ISNUMBER($F{r}),$G{r}>$F{r})," !",""),"")')
        c = G[f"{L}{r}"]
        c.font = Font(name=FONT, size=9, bold=True)
        c.alignment, c.border = center, border

grid = f"{first_week}{FIRST}:{last_week}{LAST}"
fw = first_week
G.conditional_formatting.add(grid, FormulaRule(
    formula=[f'AND({fw}{FIRST}<>"",ISNUMBER(SEARCH("!",{fw}{FIRST})))'],
    fill=PatternFill("solid", fgColor="1A2430"), font=Font(name=FONT, color="FFFFFF", bold=True),
    stopIfTrue=True))
for code, sem in (("ПКР", "ПКР"), ("РИ", "РИ"), ("УД", "УД")):
    bg, fg = SEM[sem]
    G.conditional_formatting.add(grid, FormulaRule(
        formula=[f'LEFT({fw}{FIRST},{len(code)})="{code}"'], fill=PatternFill("solid", fgColor=fg),
        font=Font(name=FONT, color="FFFFFF", bold=True), stopIfTrue=True))
for code, sem in (("ПКР", "ПКР"), ("РИ", "РИ"), ("УД", "УД")):
    bg, _fg = SEM[sem]
    G.conditional_formatting.add(grid, FormulaRule(
        formula=[f'AND($E{FIRST}="{code}",ISNUMBER($F{FIRST}),{fw}$4+6>={CALC},{fw}$4<=$F{FIRST})'],
        fill=PatternFill("solid", fgColor=bg)))
G.conditional_formatting.add(f"A{FIRST}:G{LAST}", FormulaRule(
    formula=[f'AND(ISNUMBER($D{FIRST}),$D{FIRST}<{UD_FROM})'], font=Font(name=FONT, color="9AA7B4")))
G.conditional_formatting.add(f"{first_week}4:{last_week}4", FormulaRule(
    formula=[f"AND({CALC}>={fw}$4,{CALC}<{fw}$4+7)"], fill=PatternFill("solid", fgColor=BLUE)))
G.freeze_panes = f"{first_week}{FIRST}"
G.sheet_view.zoomScale = 90

legend_row = LAST + 2
G[f"A{legend_row}"] = "Коды:"
G[f"A{legend_row}"].font = f_bold
for i, (code, text, sem) in enumerate((
        ("ПКР", "плановый комплекс мероприятий по ремонту + включение в ЭРА ремонта (документ — Промайн)", "ПКР"),
        ("РИ", "решение инженера (диапазон 61–79, мероприятие уточнить)", "РИ"),
        ("УД", "ударная дозировка", "УД"),
        ("!", "плановая дата позже крайнего срока", None))):
    r = legend_row + 1 + i
    c = G.cell(r, 2, code)
    c.alignment, c.font = center, Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=SEM[sem][1] if sem else "1A2430")
    G.cell(r, 3, text).font = f_base

# --------------------------------------------------------------------------- Сводка
S = wb.create_sheet("Сводка", 2)
S["A1"] = "Сводка по рейтингу"
S["A1"].font = f_title
S["A2"] = "Все значения считаются формулами по листу «Рейтинг»."
S["A2"].font = f_note
RQ = f"Рейтинг!${col['index']}${FIRST}:${col['index']}${LAST}"
RM = f"Рейтинг!${col['measure']}${FIRST}:${col['measure']}${LAST}"
RF = f"Рейтинг!${col['field']}${FIRST}:${col['field']}${LAST}"
RE = f"Рейтинг!${col['era']}${FIRST}:${col['era']}${LAST}"
RT = f"Рейтинг!${col['promine']}${FIRST}:${col['promine']}${LAST}"
RW = f"Рейтинг!${col['plan_status']}${FIRST}:${col['plan_status']}${LAST}"
RB = f"Рейтинг!${col['base_res']}${FIRST}:${col['base_res']}${LAST}"
RC = f"Рейтинг!${col['category']}${FIRST}:${col['category']}${LAST}"


def head_row(sheet, r, labels):
    for j, h in enumerate(labels, start=1):
        c = sheet.cell(r, j, h)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border


def body(sheet, r, values, fmt=None):
    for j, v in enumerate(values, start=1):
        c = sheet.cell(r, j, v)
        c.font, c.border = (f_base if j == 1 else f_bold), border
        c.alignment = left if j == 1 else center


head_row(S, 4, ["Мероприятие", "Скважин"])
for i, lab in enumerate((L_TO, L_RI, L_UD, L_MON)):
    body(S, 5 + i, [f"={lab}", f"=COUNTIF({RM},{lab})"])
body(S, 9, ["Всего в рейтинге", f"=COUNT({RQ})"])
S["A9"].font = f_bold

head_row(S, 11, ["Контроль исполнения", "Скважин"])
body(S, 12, ["Требуют включения в ЭРА ремонта", f'=COUNTIF({RE},"Включить")'])
body(S, 13, ["…из них документ Промайн ещё не направлен",
             f'=COUNTIFS({RE},"Включить",{RT},"<>Направлен в ЭРА")'])
body(S, 14, ["План ППР позже крайнего срока", f'=COUNTIF({RW},"Позже срока")'])
body(S, 15, ["Мероприятие назначено, но не запланировано", f'=COUNTIF({RW},"Не запланировано")'])

fields = sorted({w[1] for w in wells})
head_row(S, 17, ["Месторождение"] + [f"={x}" for x in (C_TO, C_RI, C_UD)] + ["Всего"])
for i, fld in enumerate(fields):
    r = 18 + i
    vals = [fld] + [f'=COUNTIFS({RF},$A{r},{RM},{lab})' for lab in (L_TO, L_RI, L_UD)]
    vals.append(f'=COUNTIF({RF},$A{r})')
    body(S, r, vals)

r0 = 18 + len(fields) + 1
S.cell(r0, 1, "Сверка модели с фактом разбора на базе").font = f_bold
head_row(S, r0 + 1, ["Результат разбора", "Индекс ≥ 61", "Индекс ≤ 60"])
for i, res in enumerate(("Коррозия есть", "Коррозии нет")):
    r = r0 + 2 + i
    body(S, r, [res, f'=COUNTIFS({RB},$A{r},{RQ},">"&{UD_TO})',
                f'=COUNTIFS({RB},$A{r},{RQ},"<="&{UD_TO})'])
S.cell(r0 + 4, 1, ("Коррозия есть при низком индексе или её нет при высоком — кандидаты на разбор "
                   "расхождения модели.")).font = f_note

r1 = r0 + 6
head_row(S, r1, ["Категория скважины", "Скважин"])
for i, cat in enumerate(("Коррозионная", "Среднекоррозионная", "Некоррозионная")):
    body(S, r1 + 1 + i, [cat, f'=COUNTIF({RC},$A{r1 + 1 + i})'])
S.column_dimensions["A"].width = 46
for L in "BCDE":
    S.column_dimensions[L].width = 16

# --------------------------------------------------------------------------- Справочник полей
D = wb.create_sheet("Справочник полей", 3)
D["A1"] = "Справочник полей рейтинга"
D["A1"].font = f_title
D["A2"] = ("Что хранится в каждом поле, как оно заполняется и где показывается на дашборде. "
           "«Источник» — предложение, сверить с фактическими системами.")
D["A2"].font = f_note
dheads = ["Поле", "Колонка", "Тип и значения", "Заполнение", "Источник (предложение)", "На дашборде"]
head_row(D, 4, dheads)
spec = {
    "rank": ("Целое; 1 — самый высокий индекс", "Формула", "—", "Таблица"),
    "well": ("Текст, уникален в месторождении", "Ввод", "Паспорт скважины", "Таблица"),
    "field": ("Текст", "Ввод", "Паспорт скважины", "Таблица"),
    "equip": ("Текст: тип насоса, НКТ", "Ввод", "Паспорт / ЭРА", "Карточка"),
    "status": ("В работе / Не в работе", "Ввод", "Телеметрия (DARS)", "Таблица (значок)"),
    "sampling": ("Да / Нет", "Ввод", "ХИМИК", "Карточка"),
    "env": ("Коррозионная / Некоррозионная", "Ввод", "ХИМИК", "Таблица (признаки)"),
    "tonnage": ("Число ≥ 0, т (уточнить)", "Ввод", "Служба защиты", "Карточка"),
    "noo": ("Число ≥ 0, сут (уточнить расшифровку)", "Ввод", "Уточнить", "Карточка"),
    "act_dem": ("Есть / Нет", "Ввод", "ЭРА", "Карточка"),
    "corr_sign": ("Есть / Нет / — (нет акта)", "Ввод", "ЭРА (акт демонтажа)", "Таблица (признаки)"),
    "act_base": ("Есть / Нет", "Ввод", "ЭРА", "Карточка"),
    "base_res": ("Коррозия есть / Коррозии нет / —", "Ввод", "ЭРА (акт разбора)",
                 "Таблица (признаки)"),
    "category": ("Некоррозионная / Среднекоррозионная / Коррозионная", "Ввод", "НТЦ", "Таблица"),
    "ntc": ("Число 0–1", "Модель", "Модель НТЦ", "Таблица (полоса)"),
    "math": ("Число 0–1", "Модель", "Математическая модель", "Таблица (полоса)"),
    "index": ("Целое 0–100", "Формула", "—", "Таблица"),
    "measure": ("Плановый комплекс по ремонту / Решение инженера / УД / Мероприятия не требуются",
                "Формула (в системе пересчитывается)", "—", "Таблица"),
    "era": ("Включить / —", "Формула", "—", "Таблица"),
    "promine": ("Не сформирован / Сформирован / Направлен в ЭРА / —", "Ввод", "Промайн", "Таблица"),
    "deadline": ("Дата", "Формула", "—", "Таблица, график ППР"),
    "plan": ("Дата", "Ввод", "График ППР (ДДО)", "Таблица, график ППР"),
    "plan_status": ("В срок / Позже срока / Не запланировано / —", "Формула", "—", "Таблица"),
    "comment": ("Текст", "Ввод", "Инженер", "Карточка"),
}
for i, (key, header, *_r) in enumerate(cols):
    r = 5 + i
    t, how, srcs, where = spec[key]
    body(D, r, [header, col[key], t, how, srcs, where])
    for j in range(1, 7):
        D.cell(r, j).font = f_base
        D.cell(r, j).alignment = left
for L, w in zip("ABCDEF", (26, 9, 46, 12, 26, 22)):
    D.column_dimensions[L].width = w

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {"Рейтинг": NAVY, "График ППР": BLUE}.get(ws.title, "8FA3B8")
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
wb.active = 0
wb.save(out)
print("saved", out)
